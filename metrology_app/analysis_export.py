"""Two export renderers over the already prepared data of an analysis window.

No parent Workbook inputs, fitting, or surface interpolation occur here. Map
and Radius pictures come from their existing, successfully drawn figures.
"""
from io import BytesIO
from pathlib import Path
from copy import deepcopy

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.image import Image

from .excel_export import (PLOT_COLORS, _value, _write_frame, chart_limits, plot_legend,
    plot_axis_limits, save_atomic, set_export_font, sheet_title, stage_suffix,
    style_native_series, style_plot_axis, style_workbook, sequence_layout,
    PANEL_HEIGHT_CM, MATCH_WIDTH_CM, TREND_WIDTH_CM)


def _series(name, x, y, *, line=True, markers=True, **metadata):
    return {"name": str(name), "line": line, "markers": markers, "data": pd.DataFrame({"X": x, "Y": y, **metadata})}


def capture_window(window, *, parameters=None):
    """Detach source tables and successful plot choices without changing the UI."""
    if window.refresh_timer.isActive():
        window.recognize()
    scope = str(getattr(window, "_managed_scope", ""))
    stage = next((part for part in scope.split(".") if part in ("preview", "final")), None)
    report = {"frames": {"Data": window.model.document_frame().copy(deep=True)},
              "plots": [], "pictures": [], "suffix": stage_suffix(stage) if stage else "",
              "title": window.windowTitle()}
    if window.workspace_type == "correlation_trend":
        report["frames"] = {"Ref Data": window.reference_model.document_frame().copy(deep=True),
                            "Raw Data": window.model.document_frame().copy(deep=True)}
        page = window.correlation_page
        _check_draw(page)
        if page.ready:
            for fit in page.display_fits():
                fit = page.oriented_fit(fit)
                sources = fit.sources if hasattr(fit, "sources") else (None,)
                series, notes = [], []
                for source in sources:
                    values = source.fit if source else fit
                    label = source.name if source else "Measurements"
                    series.append(_series(label, values.x.copy(), values.y.copy(), line=False))
                    order = np.argsort(values.x, kind="stable")
                    series.append(_series(f"{label} fit", values.x[order], values.predicted[order], markers=False))
                    notes.append(f"{label}: slope={values.slope:.8g}, intercept={values.intercept:.8g}, R²={values.rsquared:.8g}")
                report["plots"].append({"sheet": "Correlation", "title": f"{fit.y_name} vs {fit.x_name}",
                    "x": fit.x_name, "y": fit.y_name, "series": series, "notes": "; ".join(notes)})
        page = window.sequence_page
        _check_draw(page)
        if page.ready:
            report["plots"].extend(page.export_plot_data())
    elif window.workspace_type == "dynamic":
        page = window.dynamic_page
        for parameter, model in page.parameter_models.items():
            pivot = model.frame.copy(deep=True)
            report["frames"][f"Pivot {parameter}"] = pivot.reset_index()
            report["plots"].append({"sheet": "Dynamic", "title": f"{parameter} 3 Sigma",
                "kind": "bar", "x": "Die Seq", "y": "3 Sigma", "series": [
                    _series(parameter, list(pivot.columns), pivot.loc["3 Sigma"].to_numpy(float))]})
        page = window.dynamic_trend
        for parameter, table in page.panel_tables.items():
            die = page.selected_dies[parameter]
            report["plots"].append({"sheet": "Dynamic Trend", "title": f"{parameter} · Die {die}",
                "x": "Cycle", "y": parameter, "series": [
                    _series(f"Die {die}", np.asarray(table.index, dtype=float), table[die].to_numpy(float))]})
    else:
        for label, page in (("Wafer Map", window.plot_page), ("Radius", window.radius_page)):
            _check_draw(page)
            if getattr(page, "worker", None) is not None:
                raise ValueError("Wait for the Wafer Map drawing to finish before exporting.")
            ready = (page.result is not None and bool(page.artists)
                     if label == "Wafer Map" else page.ready)
            if not ready:
                continue
            from .appearance import MAX_EXPORT_PIXELS, export_dpi, resolution_settings
            requested = resolution_settings(page.resolution)[1]
            if parameters is None:
                buffer = BytesIO()
                page.figure.savefig(buffer, format="png", facecolor="white",
                                    dpi=export_dpi(page.figure, requested, MAX_EXPORT_PIXELS))
                report["pictures"].append((label, buffer.getvalue(),
                    f"Page {page.page_index + 1}; current software PNG"))
            else:
                # Crop existing successful software panels, never refit/interpolate
                # or alter the live figure's visibility/options to filter exports.
                from matplotlib.transforms import Bbox
                page.figure.canvas.draw()
                axes = ([(artist.axes, scene["metric"]) for artist, scene in page.artists]
                        if label == "Wafer Map" else [(axis, axis.get_title()) for axis in page.figure.axes])
                for axis, metric in axes:
                    if metric not in parameters:
                        continue
                    renderer = page.figure.canvas.get_renderer()
                    boxes = [axis.get_tightbbox(renderer)]
                    colorbar = getattr(axis.images[0], "colorbar", None) if axis.images else None
                    if colorbar is not None:
                        boxes.append(colorbar.ax.get_tightbbox(renderer))
                    bounds = Bbox.union(boxes).transformed(page.figure.dpi_scale_trans.inverted()).expanded(1.04, 1.04)
                    buffer = BytesIO()
                    page.figure.savefig(buffer, format="png", facecolor="white", bbox_inches=bounds,
                                        dpi=export_dpi(page.figure, requested, MAX_EXPORT_PIXELS))
                    report["pictures"].append((f"{label} {metric}", buffer.getvalue(), "Current software PNG"))
    return report


def _check_draw(page):
    if getattr(page, "has_drawn_once", False) and page.selector.pending_draw:
        raise ValueError("Click Draw selected to apply pending plot choices before exporting.")
    if getattr(page, "input_refresh_timer", None) is not None and page.input_refresh_timer.isActive():
        page.refresh_previous_selection()


def plot_figure(plot):
    """Render a report panel with Matplotlib, preserving all selected values."""
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    sequence = sequence_layout(plot["labels"]) if plot.get("labels") else None
    width = (sequence["width"] if sequence else MATCH_WIDTH_CM
             if plot.get("sheet") == "Correlation" or plot.get("kind") == "bar" else TREND_WIDTH_CM)
    figure = Figure(figsize=(width / 2.54, PANEL_HEIGHT_CM / 2.54), facecolor="white")
    FigureCanvasAgg(figure)
    axis = figure.subplots()
    second_axes = {}
    values_by_axis = {}
    for index, series in enumerate(plot["series"]):
        color = "#" + PLOT_COLORS[index % len(PLOT_COLORS)]
        target = axis
        if series.get("secondary"):
            key = series.get("unit", series["name"])
            if key not in second_axes:
                secondary = axis.twinx()
                secondary.spines["right"].set_position(("axes", 1 + .15 * len(second_axes)))
                secondary.set_ylabel(key)
                second_axes[key] = secondary
            target = second_axes[key]
        data = series["data"]
        x = (np.interp(data["X"].to_numpy(float), np.arange(1, len(plot["labels"]) + 1), sequence["x"])
             if sequence else data["X"])
        values_by_axis.setdefault(target, []).extend(data["Y"].to_numpy(float))
        if plot.get("kind") == "bar":
            target.bar(np.arange(len(data)), data["Y"], label=series["name"], color=color, width=.65)
            target.set_xticks(np.arange(len(data)), data["X"].astype(str))
            target.set_ylim(bottom=0)
        elif series.get("line", True):
            target.plot(x, data["Y"], marker=("o" if index % 2 == 0 else "s") if series.get("markers", True) else None,
                markersize=2.5, linewidth=1.2, linestyle="--" if index % 2 and series.get("markers", True) else "-",
                color=color, label=series["name"])
        else:
            target.scatter(x, data["Y"], s=14, color=color, edgecolors="none", label=series["name"])
    style_plot_axis(axis, title=plot["title"], xlabel=plot["x"], ylabel=plot["y"])
    for target, values in values_by_axis.items():
        plot_axis_limits(target, values, bars=plot.get("kind") == "bar")
        if target is not axis:
            target.tick_params(labelsize=10, colors="#334155")
    if plot.get("kind") != "bar" and plot["series"]:
        positions = np.concatenate([series["data"]["X"].to_numpy(float) for series in plot["series"]])
        finite = positions[np.isfinite(positions)]
        if plot.get("sheet") == "Correlation":
            plot_axis_limits(axis, values_by_axis.get(axis, []), positions)
        elif len(finite):
            padding = max(.5, (finite.max() - finite.min()) * .01)
            axis.set_xlim(finite.min() - padding, finite.max() + padding)
    if sequence:
        axis.set_xticks([position for position, _ in sequence["ticks"]],
                        [label for _, label in sequence["ticks"]], fontsize=10)
        axis.set_xlim(sequence["limits"])
    lines, labels = [], []
    for target in (axis, *second_axes.values()):
        handles, names = target.get_legend_handles_labels()
        lines.extend(handles)
        labels.extend(names)
    if lines:
        columns = max(1, int(width / 6))
        plot_legend(axis, lines, labels, columns=columns)
        legend_rows = int(np.ceil(len(labels) / columns))
        axis.title.set_position((.5, 1.))
        axis.set_title(plot["title"], pad=14 * legend_rows + 12, fontsize=16, fontweight="bold")
    else:
        legend_rows = 0
    if plot.get("notes"):
        figure.suptitle(plot["notes"].replace("; ", "\n"), fontsize=9, color="#475569")
    lines = max((len(label.splitlines()) for _, label in sequence["ticks"]), default=1) if sequence else 1
    figure.subplots_adjust(left=1.6 / width, right=1 - (.45 + .9 * len(second_axes)) / width,
                           bottom=(.9 + .45 * lines) / PANEL_HEIGHT_CM,
                           top=min(.79, 1 - (1.1 + .5 * legend_rows) / PANEL_HEIGHT_CM) if not plot.get("notes") else .70)
    set_export_font(figure)
    return figure


def append_report(book, used, report, *, native=True, images=None, prefix=""):
    """Add one window's detached tables, native charts or Matplotlib pictures."""
    suffix = report["suffix"]
    def create(title):
        sheet = book.create_sheet(sheet_title(prefix + title, used, suffix))
        sheet.freeze_panes = "A2"
        return sheet
    if "fullmap_parameters" in report:
        from .excel_export import write_fullmap
        write_fullmap(create("FullMap"), report["frames"]["Data"], report["fullmap_parameters"],
                      native=native, images=images, plot_labels=report.get("plot_labels"), pictures=report["pictures"])
        return
    for title, frame in report["frames"].items():
        _write_frame(create(title), frame)
    sheets, cursors = {}, {}
    for plot in report["plots"]:
        title = plot["sheet"]
        if title not in sheets:
            sheets[title], cursors[title] = create(title), 1
        sheet, start = sheets[title], cursors[title]
        sheet.cell(row=start, column=1, value=plot["title"])
        sheet.cell(row=start + 1, column=1, value=plot.get("notes", ""))
        chart = BarChart() if plot.get("kind") == "bar" else ScatterChart()
        chart.title, chart.x_axis.title, chart.y_axis.title = plot["title"], plot["x"], plot["y"]
        chart.width, chart.height = (MATCH_WIDTH_CM if title == "Correlation" else TREND_WIDTH_CM), PANEL_HEIGHT_CM
        row = start + 3
        secondary_axes = {}
        values_by_chart = {}
        for index, series in enumerate(plot["series"]):
            data = series["data"]
            sheet.cell(row=row, column=1, value=series["name"])
            for column, name in enumerate(data.columns, 1):
                sheet.cell(row=row + 1, column=column, value=str(name))
            first = row + 2
            for offset, values in enumerate(data.itertuples(index=False, name=None), first):
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row=offset, column=column, value=_value(value))
                    if isinstance(value, str):
                        cell.data_type = "s"
            last = first + len(data) - 1
            if len(data):
                y = Reference(sheet, min_col=2, min_row=first, max_row=last)
                if plot.get("kind") == "bar":
                    chart.add_data(Reference(sheet, min_col=2, min_row=first - 1, max_row=last), titles_from_data=True)
                    chart.set_categories(Reference(sheet, min_col=1, min_row=first, max_row=last))
                else:
                    item = Series(y, Reference(sheet, min_col=1, min_row=first, max_row=last), title=series["name"])
                    item.marker = Marker(symbol="circle", size=4)
                    if not series.get("line", True):
                        item.graphicalProperties.line.noFill = True
                    target = chart
                    if series.get("secondary"):
                        unit = series.get("unit", "Secondary axis")
                        if unit not in secondary_axes:
                            second = ScatterChart()
                            second.y_axis.axId = 200
                            second.y_axis.title = unit
                            second.y_axis.crosses = "max"
                            secondary_axes[unit] = second
                        target = secondary_axes[unit]
                    target.series.append(item)
                    style_native_series(item, index, markers=series.get("markers", True))
                values_by_chart.setdefault(target if plot.get("kind") != "bar" else chart, []).extend(data["Y"].to_numpy(float))
            row = last + 3
        for target, values in values_by_chart.items():
            chart_limits(target, values, bars=plot.get("kind") == "bar")
        if plot.get("kind") != "bar" and plot["series"]:
            chart_limits(chart, values_by_chart.get(chart, []),
                         np.concatenate([series["data"]["X"].to_numpy(float) for series in plot["series"]]))
        if native:
            if secondary_axes:
                from openpyxl.utils import get_column_letter
                for index, (unit, second) in enumerate(secondary_axes.items()):
                    panel = deepcopy(chart)
                    second.x_axis = deepcopy(panel.x_axis)
                    if len(secondary_axes) > 1:
                        panel.title = f"{plot['title']} / {unit}"
                    panel += second
                    sheet.add_chart(panel, f"{get_column_letter(8 + index * 18)}{start}")
            else:
                sheet.add_chart(chart, f"H{start}")
        else:
            figure = plot_figure(plot)
            buffer = BytesIO()
            figure.savefig(buffer, format="png", dpi=150, facecolor="white")
            buffer.seek(0)
            picture = Image(buffer)
            picture.width = figure.get_figwidth() * 96
            picture.height = picture.width * figure.get_figheight() / figure.get_figwidth()
            sheet.add_image(picture, f"H{start}")
            if images is not None:
                images.append((prefix + plot["title"] + suffix, buffer.getvalue()))
        cursors[title] = max(row, start + 25)
    for title, contents, note in report["pictures"]:
        sheet = create(title)
        sheet["A1"], sheet["A2"] = title, note
        image = Image(BytesIO(contents))
        ratio = min(1, 1200 / image.width)
        image.width, image.height = image.width * ratio, image.height * ratio
        sheet.add_image(image, "A4")
        if images is not None:
            images.append((prefix + title + suffix, contents))


def write_window(window, path, *, native=True, images=None):
    report = capture_window(window)
    book = Workbook()
    book.remove(book.active)
    append_report(book, set(), report, native=native, images=images)
    style_workbook(book)
    return save_atomic(book, Path(path).with_suffix(".xlsx"))
