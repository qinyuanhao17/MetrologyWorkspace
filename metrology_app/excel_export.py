"""Excel export: source tables, numeric plot areas and native Excel charts.

Charts are written straight into the workbook (no desktop Excel, no COM): every
series points at the numeric plot area next to its parameter block, so the file
opens with ordinary, movable Excel charts that follow cell edits. Fitted Cards
stay the software result: ``PMISH`` applies the exported slope/intercept with a
formula, and no Excel fitting function or trendline is written.
"""
from pathlib import Path
from io import BytesIO
import os
import tempfile

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.data_source import AxDataSource, StrRef, MultiLevelStrRef, MultiLevelStrData, Level, StrVal
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from openpyxl.chart.text import RichText
from openpyxl.chart.series import SeriesLabel
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties

MAX_TITLE = 31
_INVALID_TITLE = set('[]:*?/\\')
BLOCK_MIN_ROWS = 30
INFO_LABELS = (
    "Parameter", "Reference column", "Raw column", "Slope", "Intercept", "R²",
    "Valid pairs", "Bias limit", "Unit", "Card source", "Stage", "Scope", "Rows",
    "Fit",
)
SLOPE_OFFSET = INFO_LABELS.index("Slope")
INTERCEPT_OFFSET = INFO_LABELS.index("Intercept")
EXPORT_FONT = "Times New Roman"
PLOT_COLORS = ("5B9BD5", "ED7D31", "70AD47", "A879C2", "DD5965", "4FADB5")
PANEL_HEIGHT_CM = 11
MATCH_WIDTH_CM = 15
TREND_WIDTH_CM = 25


def plot_scale(values):
    """Shared data-derived outer ticks; only the axes are rounded, never values."""
    finite = np.asarray(values, dtype=float).ravel()
    finite = finite[np.isfinite(finite)]
    if not len(finite):
        return None
    low, high = float(finite.min()), float(finite.max())
    pad = (high - low) * .05 if high > low else (abs(low) * .05 or .5)
    low, high = low - pad, high + pad
    requested = (high - low) / 5
    power = 10. ** np.floor(np.log10(requested))
    step = min((1, 2, 2.5, 5, 10), key=lambda value: abs(value * power - requested)) * power
    return float(np.floor(low / step) * step), float(np.ceil(high / step) * step), float(step)


def plot_axis_limits(axis, y, x=None, *, bars=False):
    from matplotlib.ticker import MultipleLocator
    for name, values in (("y", y), ("x", x)):
        if values is None:
            continue
        scale = plot_scale(values)
        if scale:
            low, high, step = scale
            if name == "y" and bars:
                low, high = 0, max(high, step)
            getattr(axis, f"set_{name}lim")((low, high))
            getattr(axis, f"{name}axis").set_major_locator(MultipleLocator(step))


def style_native_series(series, index, *, markers=True, color=None):
    color = color or PLOT_COLORS[index % len(PLOT_COLORS)]
    series.marker = Marker(symbol="circle" if markers else None,
                           size=3 if markers else None)
    series.marker.graphicalProperties = GraphicalProperties(solidFill=color,
        ln=LineProperties(noFill=True))
    if not series.graphicalProperties.line.noFill:
        series.graphicalProperties.line = LineProperties(solidFill=color, w=19050, prstDash="solid")


def chart_limits(chart, y, x=None, *, bars=False):
    bounds = plot_scale(y)
    if bounds:
        chart.y_axis.scaling.min = 0 if bars else bounds[0]
        chart.y_axis.scaling.max = max(bounds[1], bounds[2]) if bars else bounds[1]
        chart.y_axis.majorUnit = bounds[2]
    if x is not None:
        bounds = plot_scale(x)
        if bounds:
            chart.x_axis.scaling.min, chart.x_axis.scaling.max = bounds[:2]
            chart.x_axis.majorUnit = bounds[2]


def style_plot_axis(axis, *, title, xlabel, ylabel):
    """Shared compact report frame; style only this figure, never global rcParams."""
    from matplotlib.ticker import MaxNLocator, ScalarFormatter
    axis.set_title(title, fontsize=16, fontweight="bold", pad=32)
    axis.set_xlabel(xlabel, fontsize=12, labelpad=8)
    axis.set_ylabel(ylabel, fontsize=12, labelpad=8)
    axis.set_facecolor("white")
    axis.set_axisbelow(True)
    axis.grid(True, which="major", color="#DFE3E8", linewidth=.6)
    axis.tick_params(axis="both", labelsize=10, colors="#334155", length=3, width=.6)
    for spine in axis.spines.values():
        spine.set_color("#64748B")
        spine.set_linewidth(.7)
    axis.yaxis.set_major_locator(MaxNLocator(nbins=5))
    axis.yaxis.set_major_formatter(ScalarFormatter(useOffset=False))


def plot_legend(axis, handles=None, labels=None, *, columns=3):
    if handles is None:
        handles, labels = axis.get_legend_handles_labels()
    axis.legend(handles, labels, loc="lower left", bbox_to_anchor=(0, 1.015),
                ncol=min(columns, len(labels)),
                fontsize=11, frameon=False, borderaxespad=0, handlelength=2, columnspacing=1.4)


def stage_suffix(stage):
    return " (W Card)" if str(stage).lower() == "final" else " (WO Card)"


def _value(value):
    """Convert pandas/numpy scalars to values openpyxl can write."""
    if value is None:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if not np.isfinite(value) else float(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, float) and not np.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        return value
    if isinstance(missing, bool) and missing:
        return None
    return value


def sheet_title(base, used, suffix=""):
    """Excel-legal, unique sheet name: 31 characters, no []:*?/\\."""
    text = "".join(" " if character in _INVALID_TITLE else character
                   for character in str(base)).strip() or "Sheet"
    text = text[:MAX_TITLE - len(suffix)]
    candidate, index = text + suffix, 2
    while candidate.casefold() in used:
        number = f"_{index}"
        candidate = text[:MAX_TITLE - len(number) - len(suffix)] + number + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def _write_frame(worksheet, frame):
    if frame is None or not len(frame.columns):
        return
    # Explicit coordinates: styling cells first must not shift the table down.
    for column, name in enumerate(frame.columns, start=1):
        cell = worksheet.cell(row=1, column=column, value=str(name))
        cell.data_type = "s"
        cell.font = Font(name=EXPORT_FONT, bold=True)
        samples = [len(str(name)), *(len(str(value)) for value in frame.iloc[:100, column - 1] if value is not None)]
        worksheet.column_dimensions[get_column_letter(column)].width = max(12, min(36, max(samples) + 2))
    for offset, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column, item in enumerate(row, start=1):
            cell = worksheet.cell(row=offset, column=column, value=_value(item))
            if isinstance(item, str):
                cell.data_type = "s"  # Source measurement text is never an Excel formula.


def _fill_title(worksheet):
    worksheet["A1"].font = Font(bold=True)
    worksheet.freeze_panes = worksheet["A2"]


def _label_spans(labels, limit=5):
    """Representative contiguous spans, separated by their actual plot positions."""
    runs = []
    for index, label in enumerate(labels):
        if runs and runs[-1][2] == label:
            runs[-1] = (runs[-1][0], index, label)
        else:
            runs.append((index, index, label))
    if len(runs) <= 1 or limit is None:
        return runs
    centers = np.array([(start + end) / 2 for start, end, _ in runs])
    count = min(limit, len(runs))
    indices = sorted(set(int(np.abs(centers - target).argmin())
                         for target in np.linspace(centers[0], centers[-1], count)))
    spacing = .85 * (centers[-1] - centers[0]) / (count - 1)
    selected = [indices[0]]
    for index in indices[1:]:
        if centers[index] - centers[selected[-1]] >= spacing:
            selected.append(index)
        elif index == indices[-1] and len(selected) > 1:
            selected[-1] = index
    return [runs[index] for index in selected]


def sequence_layout(labels, *, name="Wafer"):
    """Reserve measured caption widths without changing observations or row order.

    Coordinates are ordinal display positions, not Die Seq measurement values.
    PNG spacing is locally expanded; ordinary Excel category charts instead
    receive enough physical width to keep the same captions readable.
    """
    from matplotlib.backends.backend_agg import RendererAgg
    from matplotlib.font_manager import FontProperties
    from .plotting.sequence_axis import WaferSpacing
    labels = list(labels)
    count = len(labels)
    limits = (.5, max(1, count) + .5)
    if not count:
        return {"name": name, "x": np.array([]), "ticks": [], "width": TREND_WIDTH_CM,
                "native_width": TREND_WIDTH_CM, "limits": limits}
    renderer, font = RendererAgg(1, 1, 96), FontProperties(family=EXPORT_FONT, size=11)
    measure = lambda text: renderer.get_text_width_height_descent(text, font, False)[0]
    runs = _label_spans(labels, limit=None)
    # A dense single Wafer still fits the 25 cm base panel. Only captions that
    # genuinely need more room add width, rather than two pixels per source row.
    factor = min(1., (TREND_WIDTH_CM * 96 / 2.54 - 110) / (2 * count))
    spans = [((start + .5) * factor, (end + 1.5) * factor, label) for start, end, label in runs]
    spacing = WaferSpacing(spans, tuple(value * factor for value in limits), measure)
    ticks = [(float(spacing.forward((start + end + 2) / 2 * factor)) / factor, label)
             for start, end, label in runs]
    centers = np.array([(start + end + 2) / 2 for start, end, _ in runs])
    widths = np.array([max((measure(line) for line in text.splitlines()), default=0) + 16 for _, _, text in runs])
    required = max(widths[0] * count / (2 * (centers[0] - .5)),
                   widths[-1] * count / (2 * (count + .5 - centers[-1])))
    if len(runs) > 1:
        required = max(required, float(np.max((widths[1:] + widths[:-1]) / 2 * count / np.diff(centers))))
    return {"name": name, "x": spacing.forward(np.arange(1, count + 1) * factor) / factor,
            "native_width": max(TREND_WIDTH_CM, (required + 110) * 2.54 / 96),
            "ticks": ticks, "width": max(TREND_WIDTH_CM, (spacing.width + 110) * 2.54 / 96), "limits": limits}


def native_categories(sheet, first, last, identities):
    """Ordinary local multi-level categories, editable with Excel range handles."""
    levels = []
    for column, values in reversed(identities):
        points = [StrVal(idx=index, v=str(value)) for index, value in enumerate(values)
                  if index == 0 or (len(identities) > 1 and column == identities[-1][0]) or value != values[index - 1]]
        levels.append(Level(pt=points))
    return AxDataSource(multiLvlStrRef=MultiLevelStrRef(
        f=str(Reference(sheet, min_col=identities[0][0], max_col=identities[-1][0], min_row=first, max_row=last)),
        # ptCount is optional in OOXML; openpyxl 3.1 cannot serialize its
        # plain-integer descriptor here. The complete inner level sets depth.
        multiLvlStrCache=MultiLevelStrData(lvl=levels)))


def _next_plot_column(sheet, column, pixels):
    """Leave a gutter after a drawing, accounting for this sheet's real widths."""
    pixels += 24
    while pixels > 0:
        if column > 16384:
            raise ValueError("Repeated identity captions exceed Excel's worksheet width. Collect rows by Wafer before exporting.")
        width = sheet.column_dimensions[get_column_letter(column)].width
        pixels -= (width or 13) * 7 + 5
        column += 1
    return column


def _original_value(original, column, source_row):
    source = f"'{original.title}'!${column}${int(source_row) + 2}"
    return f'=IF({source}="",NA(),IFERROR({source}*1,NA()))'


def _write_parameter_block(worksheet, row, parameter, mapping, data, card,
                           card_source, scope, stage, *, native=True, images=None, show_percent=False, reference_label="REF", original=None, wafer_only=False):
    """Compact visible measurements; chart labels/fit provenance stay internal."""
    reference = data["Reference"].to_numpy(float)
    raw = data["Raw"].to_numpy(float)
    count = len(data)
    slope = float(card.slope) if card is not None and np.isfinite(card.slope) else None
    intercept = (float(card.intercept)
                 if card is not None and np.isfinite(card.intercept) else None)
    equation = ""
    if slope is not None and intercept is not None:
        sign = "+" if intercept >= 0 else "-"
        equation = f"y = {slope:.6g}x {sign} {abs(intercept):.6g}"
        if card is not None and np.isfinite(card.r_squared):
            equation += f" · R² = {card.r_squared:.6g}"
    facts = (
        parameter,
        "" if mapping is None else mapping.reference_column,
        parameter if mapping is None else mapping.raw_column,
        slope,
        intercept,
        None if card is None or not np.isfinite(card.r_squared) else float(card.r_squared),
        None if card is None else int(card.valid_pairs),
        "" if mapping is None or mapping.bias_limit is None else float(mapping.bias_limit),
        "" if mapping is None else _unit(mapping.reference_column),
        card_source,
        stage,
        scope,
        count,
        equation,
    )
    helper = worksheet.parent["_Plot data"]
    helper_row = helper.max_row + 3
    for column, text in enumerate(("Source row", "Axis label", "Fit X", "Fit Y"), 1):
        helper.cell(row=helper_row - 1, column=column, value=text)
    for offset, (label, value) in enumerate(zip(INFO_LABELS, facts)):
        helper.cell(row=helper_row + offset, column=10, value=label)
        helper.cell(row=helper_row + offset, column=11, value=_value(value))
    slope_cell = f"'{helper.title}'!$K${helper_row + SLOPE_OFFSET}"
    intercept_cell = f"'{helper.title}'!$K${helper_row + INTERCEPT_OFFSET}"
    final = stage.lower() == "final"
    identities = (["Group"] if "Group label" in data else []) + ["Wafer", "Die Seq"]
    headers = identities + ["PMISH"] + ([] if final else ["PMISH+Card"]) + [reference_label, "Bias", "Bias %"]
    columns = {name: index for index, name in enumerate(headers, 1)}
    worksheet.merge_cells(start_row=row, start_column=1,
                          end_row=row, end_column=len(headers))
    title = worksheet.cell(row=row, column=1, value=parameter)
    title.font = Font(name=EXPORT_FONT, size=14, bold=True)
    from openpyxl.styles import Alignment
    title.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[row].height = 22
    worksheet.row_dimensions[row + 1].height = 20
    header_row = row + 1
    for column, text in enumerate(headers, start=1):
        cell = worksheet.cell(row=header_row, column=column, value=text)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF2F7")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        letter = get_column_letter(column)
        worksheet.column_dimensions[letter].width = max(worksheet.column_dimensions[letter].width,
                                                        30 if text == "Wafer" else 16 if text == "Group" else 14)
    first_row = header_row + 1
    fit = slope is not None and intercept is not None
    primary_labels = data["Wafer"] if wafer_only else data.get("Group label", data.get("Wafer", data.get("Label", pd.Series([""] * count))))
    axis_labels = {(start + end) // 2: label for start, end, label in _label_spans(primary_labels, limit=None)}
    source_column = (next(cell.column_letter for cell in original[1] if cell.value == mapping.raw_column)
                     if original is not None and mapping is not None else None)
    reference_column = (next(cell.column_letter for cell in reversed(original[1])
                             if cell.value == f"{parameter} Reference ({reference_label})")
                        if original is not None else None)
    for index in range(count):
        target = first_row + index
        for name in identities:
            source = "Group label" if name == "Group" else name
            value = data[source].iloc[index] if source in data else ""
            cell = worksheet.cell(row=target, column=columns[name], value=_value(value))
            if isinstance(value, str):
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        worksheet.row_dimensions[target].height = 15
        raw_cell = f"{get_column_letter(columns['PMISH'])}{target}"
        raw_value = _value(raw[index])
        if source_column is not None:
            raw_value = _original_value(original, source_column, data["Source row"].iloc[index])
        worksheet.cell(row=target, column=columns["PMISH"], value=raw_value)
        # Final Raw Data is already the accepted post-Card source. Applying
        # the fitted Card here a second time would disagree with the software.
        if not final:
            worksheet.cell(row=target, column=columns["PMISH+Card"],
                           value=f'=IF(ISNUMBER({raw_cell}),{slope_cell}*{raw_cell}+{intercept_cell},NA())' if fit else None)
        evaluated_cell = f"{get_column_letter(columns['PMISH' if final else 'PMISH+Card'])}{target}"
        ref_cell = f"{get_column_letter(columns[reference_label])}{target}"
        worksheet.cell(row=target, column=columns[reference_label],
                       value=_original_value(original, reference_column, data["Source row"].iloc[index])
                       if reference_column is not None else _value(reference[index]))
        worksheet.cell(row=target, column=columns["Bias"],
                       value=f'=IF(AND(ISNUMBER({evaluated_cell}),ISNUMBER({ref_cell})),{evaluated_cell}-{ref_cell},NA())')
        worksheet.cell(row=target, column=columns["Bias %"],
                       value=(f'=IF(AND(ISNUMBER({evaluated_cell}),ISNUMBER({ref_cell})),'
                              f'IF({ref_cell}=0,NA(),({evaluated_cell}-{ref_cell})/{ref_cell}*100),NA())'))
        for column in range(len(identities) + 1, len(headers) + 1):
            worksheet.cell(row=target, column=column).alignment = Alignment(vertical="center")
        source_row = data["Source row"].iloc[index] if "Source row" in data else data.index[index]
        helper.cell(row=helper_row + index, column=1, value=int(source_row) + 1)
        cell = helper.cell(row=helper_row + index, column=2, value=axis_labels.get(index, ""))
        cell.data_type = "s"
    # Merge only the derived presentation column. A hidden Lot/PAD caption must
    # never merge different measurement identities, or cross a Group boundary.
    keys = [data["Wafer key"].iloc[index] if wafer_only else (data["Group"].iloc[index], data["Wafer key"].iloc[index])
            for index in range(count)]
    for start, end, _key in _label_spans(keys, limit=None):
        if end > start:
            worksheet.merge_cells(start_row=first_row + start, end_row=first_row + end,
                                  start_column=columns["Wafer"], end_column=columns["Wafer"])
        cell = worksheet.cell(first_row + start, columns["Wafer"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if end == start:
            worksheet.row_dimensions[first_row + start].height = 15 * max(1, len(str(cell.value).splitlines()))
    if fit:
        finite = raw[np.isfinite(raw) & np.isfinite(reference)]
        if len(finite):
            for offset, value in enumerate((float(np.min(finite)), float(np.max(finite)))):
                helper.cell(row=helper_row + offset, column=3, value=value)
                helper.cell(row=helper_row + offset, column=4,
                            value=f"={slope_cell}*C{helper_row + offset}+{intercept_cell}")
    last_row = first_row + count - 1
    evaluated_values = raw if stage.lower() == "final" else (slope * raw + intercept if fit else np.full(count, np.nan))
    percentages = np.full(count, np.nan)
    np.divide((evaluated_values - reference) * 100, reference, out=percentages,
              where=np.isfinite(reference) & (reference != 0))
    layouts = [sequence_layout(primary_labels, name="Group" if "Group label" in data and not wafer_only else "Wafer")]
    if "Group label" in data and not wafer_only:
        labels = [f"{group}\n{wafer}" for group, wafer in zip(data["Group label"], data["Wafer"])]
        layouts.append(sequence_layout(labels))
    category_columns = [(columns["Group"], data["Group label"].tolist())] if "Group" in columns and not wafer_only else None
    detailed_columns = ([(columns["Group"], data["Group label"].tolist())] if "Group" in columns and not wafer_only else []) + [
        (columns["Wafer"], data["Wafer"].tolist()), (columns["Die Seq"], data["Die Seq"].tolist())]
    for index, layout in enumerate(layouts):
        layout["categories"] = native_categories(worksheet, first_row, last_row,
            category_columns if index == 0 and category_columns else detailed_columns) if count else None
    if count and native:
        _add_charts(worksheet, row, first_row, last_row, parameter, fit, show_percent=show_percent,
                    raw=raw, reference=reference, evaluated=evaluated_values, equation=equation,
                    reference_label=reference_label, unit="" if mapping is None else _unit(mapping.reference_column),
                    columns=columns, helper=helper, helper_row=helper_row, final=final, layouts=layouts)
    elif count:
        anchor = len(headers) + 2
        panels = [("Match", None)] + [(panel, layout) for layout in layouts
            for panel in ("Trend", "Bias", *(("Bias %",) if show_percent else ()))]
        for panel, layout in panels:
            figure = parameter_figure(parameter, data, card, stage, scope, panel=panel,
                reference_label=reference_label, unit="" if mapping is None else _unit(mapping.reference_column), sequence=layout)
            _embed_figure(worksheet, figure, f"{get_column_letter(anchor)}{row}",
                          f"{parameter} {panel} {layout['name'] if layout else ''} {scope} {stage}{stage_suffix(stage)}", images)
            anchor = _next_plot_column(worksheet, anchor, figure.get_figwidth() * 96)
            figure.clear()
    return max(first_row + count, row + BLOCK_MIN_ROWS) + 3


def _add_charts(worksheet, row, first_row, last_row, parameter, fit, *, show_percent=False,
                raw, reference, evaluated, equation, reference_label, unit, columns, helper, helper_row, final, layouts):
    anchor = len(columns) + 2
    def add(chart):
        nonlocal anchor
        chart.visible_cells_only = False  # Fit endpoints/captions are on the hidden plot-data sheet.
        worksheet.add_chart(chart, f"{get_column_letter(anchor)}{row}")
        anchor = _next_plot_column(worksheet, anchor, chart.width * 96 / 2.54)
    match = ScatterChart()
    match.title = parameter + ("\n" + equation.replace(" · ", "\n") if equation else "")
    match.x_axis.title = "PMISH"
    match.y_axis.title = reference_label
    match.height, match.width = PANEL_HEIGHT_CM, MATCH_WIDTH_CM
    match.legend = None
    match.x_axis.delete = False
    match.y_axis.delete = False
    points = Series(Reference(worksheet, min_col=columns[reference_label], min_row=first_row, max_row=last_row),
                    Reference(worksheet, min_col=columns["PMISH"], min_row=first_row, max_row=last_row),
                    title=reference_label)
    points.marker = Marker(symbol="circle", size=5)
    points.graphicalProperties.line.noFill = True
    style_native_series(points, 0, color="4F8BD6")
    match.series.append(points)
    if fit:
        fit_series = Series(
            Reference(helper, min_col=4, min_row=helper_row, max_row=helper_row + 1),
            Reference(helper, min_col=3, min_row=helper_row, max_row=helper_row + 1),
            title="Fit")
        fit_series.marker = Marker(symbol=None)
        style_native_series(fit_series, 1, markers=False, color="E09F3E")
        match.series.append(fit_series)
    add(match)
    valid = np.isfinite(raw) & np.isfinite(reference)
    chart_limits(match, reference[valid], raw[valid])

    percentages = np.full(len(reference), np.nan)
    np.divide((evaluated - reference) * 100, reference, out=percentages,
              where=np.isfinite(reference) & (reference != 0))
    for layout in layouts:
        panels = [("Trend", ("PMISH" if final else "PMISH+Card", reference_label), np.concatenate((evaluated, reference)), 3),
                  ("Bias", ("Bias",), evaluated - reference, 4)]
        if show_percent:
            panels.append(("Bias %", ("Bias %",), percentages, 5))
        for panel, names, values, _baseline_column in panels:
            chart = LineChart()
            chart._identity_count = last_row - first_row + 1
            chart.title = parameter if panel == "Trend" else panel
            chart.y_axis.title = parameter if panel == "Trend" else f"Bias ({unit})" if panel == "Bias" and unit and unit != "1" else panel
            chart.height, chart.width = PANEL_HEIGHT_CM, layout["native_width"]
            chart.x_axis.title = layout["name"]
            if panel != "Trend":
                chart.legend = None
            for name in names:
                chart.add_data(Reference(worksheet, min_col=columns[name], min_row=first_row - 1, max_row=last_row), titles_from_data=True)
                chart.series[-1].cat = layout["categories"]
                if name == reference_label:
                    style_native_series(chart.series[-1], 1, color="ED7D31")
            chart_limits(chart, values)
            add(chart)


def _unit(reference_column):
    from .trend import parse_unit
    return parse_unit(reference_column)


def _parameter_data(result, parameter, key="all", labels=None, *, order="group"):
    from .match_groups import normalized, wafer_identity_keys
    labels = labels or {"pad": True, "lot": True}
    data = result.group_series(parameter, key, trend=True, card_mode="group" if key != "all" else "raw")
    raw = result.group_plan.raw
    identities = wafer_identity_keys(raw)
    die_column = next((column for column in raw if normalized(column) in ("dieseq", "dieid", "diesequence")), None)
    captions, dies, wafer_captions = [], [], []
    for row, group in zip(data["Source row"], data["Group"]):
        wafer, lot, pad = identities[row] if identities is not None else (f"Row {row + 1}", "", "")
        text = [wafer]
        if lot and labels.get("lot"):
            text.append(f"LOT: {lot}")
        if pad and labels.get("pad"):
            text.append(f"PAD: {pad}")
        wafer_captions.append("\n".join(text))
        if result.group_plan.enabled:
            group_label = result.group_plan.label(group)
            text = ([group_label, *text] if group in result.group_plan.state.get("show_wafer_groups", ()) else [group_label])
        captions.append("\n".join(text))
        dies.append(str(raw.iloc[row][die_column]) if die_column is not None else "")
    data["Label"], data["Die Seq"] = captions, dies
    data["Wafer"] = wafer_captions
    data["Wafer key"] = [identities[row] if identities is not None else (f"Row {row + 1}", "", "")
                         for row in data["Source row"]]
    if result.group_plan.enabled:
        data["Group label"] = data["Group"].map(result.group_plan.label)
    if order == "wafer":
        # First appearance in the Original source, independent of Group order
        # or the UI's selected trend order. Keep original rows within a Wafer.
        wafer_ranks = {}
        for identity in identities if identities is not None else data["Wafer key"]:
            wafer_ranks.setdefault(identity, len(wafer_ranks))
        ranks = [(wafer_ranks[identity], int(row)) for identity, row in zip(data["Wafer key"], data["Source row"])]
        data = data.iloc[sorted(range(len(data)), key=ranks.__getitem__)].copy()
    elif result.group_plan.enabled:
        # Only the grouped export projection is collected. Source row IDs stay
        # attached to every observation and Original retains the untouched order.
        group_ranks = {key: index for index, key in enumerate(result.group_plan.group_keys)}
        measurement_ranks = {}
        ranks = []
        for group, identity in zip(data["Group"], data["Wafer key"]):
            group_ranks.setdefault(group, len(group_ranks))
            key = group, identity
            measurement_ranks.setdefault(key, len(measurement_ranks))
            ranks.append((group_ranks[group], measurement_ranks[key]))
        data = data.iloc[sorted(range(len(data)), key=ranks.__getitem__)].copy()
    return data


def write_parameter_sheets(book, used_titles, result, workbook, *, native=True, images=None, plot_labels=None, raw=None):
    """Original data, Match and Wafer projections, plus one sheet per Group."""
    stage = result.result_mode.title()
    suffix = stage_suffix(stage)
    original = book.create_sheet(sheet_title("Original", used_titles, suffix))
    _fill_title(original)
    original_raw = (raw if raw is not None else result.group_plan.raw).reset_index(drop=True)
    paired = pd.DataFrame(index=range(len(original_raw)))
    for parameter in result.parameter_names:
        values = np.full(len(original_raw), np.nan)
        data = result.group_series(parameter, "all", card_mode="raw")
        values[data["Source row"].to_numpy(int)] = data["Reference"].to_numpy(float)
        paired[f"{parameter} Reference ({workbook.match_type})"] = values
    # Keep the literal Raw Data columns intact; matched Reference values occupy
    # a separate area to their right so every measurement series can trace here.
    _write_frame(original, pd.concat([original_raw, pd.DataFrame({"": [None] * len(original_raw)}), paired], axis=1))
    sheet = book.create_sheet(sheet_title("Match", used_titles, suffix))
    _fill_title(sheet)
    _write_frame(sheet, raw)
    row = len(raw) + 4 if raw is not None else 1
    for parameter in result.parameter_names:
        mapping = next((item for item in workbook.mappings if item.name == parameter), None)
        row = _write_parameter_block(
            sheet, row, parameter, mapping, _parameter_data(result, parameter, labels=plot_labels),
            result.card(parameter), "Overall", "All", stage, native=native, images=images,
            show_percent="percent" in workbook.bias_views, reference_label=workbook.match_type, original=original)
    wafer = book.create_sheet(sheet_title("Wafer", used_titles, suffix))
    _fill_title(wafer)
    row = 1
    for parameter in result.parameter_names:
        mapping = next((item for item in workbook.mappings if item.name == parameter), None)
        row = _write_parameter_block(
            wafer, row, parameter, mapping, _parameter_data(result, parameter, labels=plot_labels, order="wafer"),
            result.card(parameter), "Overall", "Wafer", stage, native=native, images=images,
            show_percent="percent" in workbook.bias_views, reference_label=workbook.match_type,
            original=original, wafer_only=True)
    plan = result.group_plan
    if not plan.enabled:
        return sheet
    for key in plan.group_keys:
        label = plan.label(key)
        sheet = book.create_sheet(sheet_title(f"Match {label}", used_titles, suffix))
        _fill_title(sheet)
        row = 1
        for parameter in result.parameter_names:
            mapping = next((item for item in workbook.mappings if item.name == parameter), None)
            card = result.group_card(parameter, key)
            row = _write_parameter_block(
                sheet, row, parameter, mapping,
                _parameter_data(result, parameter, key, plot_labels),
                card if card is not None else result.card(parameter),
                "Group" if card is not None else "Overall", label, stage,
                native=native, images=images, show_percent="percent" in workbook.bias_views,
                reference_label=workbook.match_type, original=original)
    return sheet


def _wafer_label(text):
    return " / ".join(str(text).splitlines())


def write_rsq_sheet(book, used_titles, result, *, native=True, images=None):
    """One row per Wafer, one R² column per parameter, plus a chart per parameter."""
    summaries = {}
    for parameter in result.parameter_names:
        try:
            summary = result.wafer_summary(parameter)
        except ValueError:
            summary = pd.DataFrame()
        if len(summary):
            summaries[parameter] = summary
    if not summaries:
        return None
    order, groups = [], {}
    for parameter, summary in summaries.items():
        for wafer, group in zip(summary["Wafer"], summary["Group"]):
            label = _wafer_label(wafer)
            if label not in groups:
                order.append(label)
                groups[label] = group
    sheet = book.create_sheet(sheet_title("RSQ", used_titles, stage_suffix(result.result_mode)))
    _fill_title(sheet)
    sheet.cell(row=1, column=1, value="Wafer")
    sheet.cell(row=1, column=2, value="Group")
    for index, parameter in enumerate(summaries, start=3):
        cell = sheet.cell(row=1, column=index, value=f"{parameter} R²")
        cell.font = Font(bold=True)
    row_of = {}
    for offset, label in enumerate(order, start=2):
        row_of[label] = offset
        sheet.cell(row=offset, column=1, value=label)
        sheet.cell(row=offset, column=2, value=groups[label])
    for index, (parameter, summary) in enumerate(summaries.items(), start=3):
        for wafer, r_squared in zip(summary["Wafer"], summary["R²"]):
            row = row_of[_wafer_label(wafer)]
            sheet.cell(row=row, column=index, value=_value(r_squared))
    last = max(row_of.values()) if row_of else 1
    anchor_row = 1
    anchor_column = get_column_letter(len(summaries) + 4)
    for index, parameter in enumerate(summaries, start=3) if native else ():
        chart = LineChart()
        chart.title = f"{parameter} R² by Wafer"
        chart.y_axis.title = "R²"
        chart.height, chart.width = 8.0, 16.0
        chart.add_data(Reference(sheet, min_col=index, min_row=1, max_row=last),
                       titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=last))
        chart_limits(chart, [sheet.cell(row=row, column=index).value for row in range(2, last + 1)])
        sheet.add_chart(chart, f"{anchor_column}{anchor_row + (index - 3) * 25}")
    if not native:
        from .analysis_export import _series, plot_figure
        for index, (parameter, summary) in enumerate(summaries.items()):
            plot = {"title": f"{parameter} R² by Wafer", "x": "Wafer", "y": "R²",
                    "series": [_series(parameter, np.arange(1, len(summary) + 1), summary["R²"].to_numpy(float))],
                    "labels": summary["Wafer"].map(_wafer_label).tolist()}
            _embed_figure(sheet, plot_figure(plot), f"{anchor_column}{anchor_row + index * 25}",
                          plot["title"] + stage_suffix(result.result_mode), images)
    return sheet


def write_all_trend_sheet(book, used_titles, result, workbook, *, native=True, images=None):
    """Per parameter: each Group keeps its own ordinals, PMISH, REF and Bias."""
    plan = result.group_plan
    if not plan.enabled:
        return None
    sheet = book.create_sheet(sheet_title("All Trend", used_titles, stage_suffix(result.result_mode)))
    original = book["Original" + stage_suffix(result.result_mode)]
    _fill_title(sheet)
    row = 1
    for parameter in result.parameter_names:
        series = []
        mapping = next(item for item in workbook.mappings if item.name == parameter)
        source_column = next(cell.column_letter for cell in original[1] if cell.value == mapping.raw_column)
        reference_column = next(cell.column_letter for cell in reversed(original[1])
                                if cell.value == f"{parameter} Reference ({workbook.match_type})")
        for key in plan.group_keys:
            card = result.group_card(parameter, key)
            series.append((plan.label(key), key,
                           _parameter_data(result, parameter, key),
                           card if card is not None else result.card(parameter)))
        data_base = 2
        info_base = data_base + 4 * len(series) + 1
        sheet.cell(row=row, column=1, value="Parameter").font = Font(bold=True)
        sheet.cell(row=row, column=2, value=parameter)
        for slot, (label, _key, _data, card) in enumerate(series):
            column = info_base + slot * 2
            sheet.cell(row=row, column=column, value=f"{label} slope").font = Font(bold=True)
            sheet.cell(row=row, column=column + 1,
                       value=None if card is None or not np.isfinite(card.slope)
                       else float(card.slope))
            sheet.cell(row=row + 1, column=column,
                       value=f"{label} intercept").font = Font(bold=True)
            sheet.cell(row=row + 1, column=column + 1,
                       value=None if card is None or not np.isfinite(card.intercept)
                       else float(card.intercept))
        header_row = row + 2
        first_row = header_row + 1
        sheet.cell(row=header_row, column=1, value="Row").font = Font(bold=True)
        for slot, (label, _key, _data, _card) in enumerate(series):
            base = data_base + slot * 4
            for offset, name in enumerate(("Raw value", "PMISH", "REF", "Bias")):
                cell = sheet.cell(row=header_row, column=base + offset,
                                  value=f"{label} {name}")
                cell.font = Font(bold=True)
        depth = max((len(data) for _label, _key, data, _card in series), default=0)
        for index in range(depth):
            target = first_row + index
            sheet.cell(row=target, column=1, value=index + 1)
            for slot, (label, _key, data, _card) in enumerate(series):
                if index >= len(data):
                    continue
                base = data_base + slot * 4
                slope_cell = f"${get_column_letter(info_base + slot * 2 + 1)}${row}"
                intercept_cell = f"${get_column_letter(info_base + slot * 2 + 1)}${row + 1}"
                source_row = data["Source row"].iloc[index]
                sheet.cell(row=target, column=base, value=_original_value(original, source_column, source_row))
                raw_cell = f"{get_column_letter(base)}{target}"
                evaluated = (raw_cell if result.result_mode == "final" else
                             f"{slope_cell}*{raw_cell}+{intercept_cell}")
                sheet.cell(row=target, column=base + 1, value=f'=IF(ISNUMBER({raw_cell}),{evaluated},NA())')
                sheet.cell(row=target, column=base + 2, value=_original_value(original, reference_column, source_row))
                pmish_cell, ref_cell = f"{get_column_letter(base + 1)}{target}", f"{get_column_letter(base + 2)}{target}"
                sheet.cell(row=target, column=base + 3,
                           value=f'=IF(AND(ISNUMBER({pmish_cell}),ISNUMBER({ref_cell})),{pmish_cell}-{ref_cell},NA())')
        last_row = first_row + depth - 1
        if depth and native:
            chart = LineChart()
            chart.title = parameter
            chart.y_axis.title = parameter
            chart.height, chart.width = PANEL_HEIGHT_CM, TREND_WIDTH_CM
            for slot, (label, _key, data, _card) in enumerate(series):
                if not len(data):
                    continue
                base = data_base + slot * 4
                for offset in (1, 2):
                    chart.add_data(
                        Reference(sheet, min_col=base + offset, min_row=header_row,
                                  max_row=first_row + len(data) - 1),
                        titles_from_data=True)
                    chart.series[-1].tx = SeriesLabel(v=f"{label} {'PMISH' if offset == 1 else workbook.match_type}")
            chart_limits(chart, np.concatenate([
                array for _label, _key, data, card in series
                for array in ((data["Raw"].to_numpy(float) if result.result_mode == "final" else
                               card.slope * data["Raw"].to_numpy(float) + card.intercept), data["Reference"].to_numpy(float))]))
            chart.set_categories(
                Reference(sheet, min_col=1, min_row=first_row, max_row=last_row))
            sheet.add_chart(chart, f"{get_column_letter(info_base + len(series) * 2 + 2)}{row}")
        elif depth:
            from .analysis_export import _series, plot_figure
            curves = []
            for label, _key, data, card in series:
                raw = data["Raw"].to_numpy(float)
                evaluated = raw if result.result_mode == "final" else card.slope * raw + card.intercept
                curves.extend([_series(f"{label} PMISH", np.arange(1, len(data) + 1), evaluated),
                               _series(f"{label} {workbook.match_type}", np.arange(1, len(data) + 1), data["Reference"].to_numpy(float))])
            plot = {"title": f"{parameter} Group trends", "x": "Row", "y": parameter, "series": curves}
            _embed_figure(sheet, plot_figure(plot), f"{get_column_letter(info_base + len(series) * 2 + 2)}{row}",
                          plot["title"] + stage_suffix(result.result_mode), images)
        row = first_row + max(depth, 18) + 22
    return sheet


def write_workbook(path, *, reference, stages, native=True, images=None, reports=(), extra_frames=(), plot_labels=None, fullmap_parameters=(), cancelled=lambda: False):
    """Write the workbook export: data sheets plus native-chart Match sheets."""
    book = Workbook()
    book.remove(book.active)
    used = set()

    def data_sheet(title, frame):
        sheet = book.create_sheet(sheet_title(title, used))
        _fill_title(sheet)
        _write_frame(sheet, frame)
        return sheet

    helper = book.create_sheet(sheet_title("_Plot data", used))
    helper.sheet_state = "hidden"
    _write_frame(helper, reference)  # Preserve original Reference names and measurement strings.
    for workbook, result, raw, stage_frames in stages:
        if cancelled():
            raise InterruptedError("Export cancelled")
        suffix = stage_suffix(workbook.result_mode)
        write_parameter_sheets(book, used, result, workbook, native=native, images=images, plot_labels=plot_labels, raw=raw)
        write_all_trend_sheet(book, used, result, workbook, native=native, images=images)
        if str(workbook.match_type).upper() in {"NOVA", "KLA"}:
            write_rsq_sheet(book, used, result, native=native, images=images)
        if result.group_plan.enabled:
            fits = pd.concat([result.group_summary(p).assign(Parameter=p)
                              for p in result.parameter_names], ignore_index=True)
            data_sheet(sheet_title("Group Fits", set(), suffix), fits)
            data_sheet(sheet_title("Order", set(), suffix), result.group_plan.order_frame)
        fullmap_report = next((report for _prefix, report in reports
            if report.get("scope") == f"map.{workbook.result_mode}"), None)
        frames = list(stage_frames)
        if fullmap_report and not any(name == "FullMap" for name, _frame in frames):
            frames.append(("FullMap", fullmap_report["frames"]["Data"]))
        for name, frame in frames:
            if name == "FullMap":
                sheet = book.create_sheet(sheet_title(name, used, suffix))
                _fill_title(sheet)
                write_fullmap(sheet, fullmap_report["frames"]["Data"] if fullmap_report else frame,
                              [mapping.name for mapping in workbook.mappings], native=native, images=images,
                              plot_labels=plot_labels, pictures=fullmap_report["pictures"] if fullmap_report else ())
            else:
                data_sheet(sheet_title(name, set(), suffix), frame)
    for name, frame in extra_frames:
        if name.startswith("FullMap"):
            write_fullmap(book.create_sheet(sheet_title(name, used)), frame, fullmap_parameters,
                          native=native, images=images, plot_labels=plot_labels)
        else:
            data_sheet(name, frame)
    if reports:
        from .analysis_export import append_report
        for prefix, report in reports:
            if report.get("scope") in {f"map.{workbook.result_mode}" for workbook, _result, _raw, _frames in stages}:
                continue
            append_report(book, used, report, prefix=prefix, native=native, images=images)
    style_workbook(book)
    book.active = next(index for index, sheet in enumerate(book) if sheet.sheet_state == "visible")
    return save_atomic(book, path, cancelled=cancelled)


def style_workbook(book):
    """Use the Excel theme for charts, with one explicit default typeface."""
    from copy import copy
    properties = CharacterProperties(latin=DrawingFont(typeface=EXPORT_FONT), sz=1400)
    def chart_text():
        tick_properties = CharacterProperties(latin=DrawingFont(typeface=EXPORT_FONT), sz=1100)
        return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=tick_properties), endParaRPr=tick_properties)])
    # Font IDs are shared by every cell in this newly authored workbook.
    # Change their definitions once, retaining bold/size/colour. Iterating a
    # large worksheet also materializes all blank cells between plot blocks.
    for index, original in enumerate(book._fonts):
        font = copy(original)
        font.name = EXPORT_FONT
        font.scheme = None
        book._fonts[index] = font
    for sheet in book:
        for chart in [child for root in sheet._charts for child in root._charts]:
            chart.style = 2
            chart.varyColors = False
            if hasattr(chart, "smooth"):
                chart.smooth = False
            for series in chart.series:
                series.smooth = False
            for index, series in enumerate(chart.series):
                if (series.graphicalProperties.line.solidFill is None
                        and (series.marker is None or series.marker.graphicalProperties.solidFill is None)):
                    style_native_series(series, index)
                if getattr(chart, "type", None) in ("col", "bar"):
                    series.graphicalProperties.solidFill = PLOT_COLORS[index % len(PLOT_COLORS)]
            chart.x_axis.axPos = "b"
            chart.y_axis.axPos = "r" if chart.y_axis.axId == 200 else "l"
            chart.x_axis.delete = chart.y_axis.delete = False
            chart.roundedCorners = False
            chart.x_axis.tickLblPos = "low"
            chart.y_axis.tickLblPos = "nextTo"
            if hasattr(chart.x_axis, "noMultiLvlLbl"):
                chart.x_axis.noMultiLvlLbl = False
                if any(series.cat is not None and series.cat.multiLvlStrRef is not None for series in chart.series):
                    count = getattr(chart, "_identity_count", 1)
                    chart.x_axis.tickLblSkip = max(1, int(np.ceil(count * 35 / (chart.width * 96 / 2.54))))
                    chart.x_axis.majorTickMark = "none"
            # Use Excel's own chart theme rather than copying Matplotlib's frame.
            chart.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(
                ln=LineProperties(solidFill="D9D9D9", w=9525)))
            chart.x_axis.majorGridlines = None
            chart.display_blanks = "gap"
            chart.txPr = chart_text()
            if chart.legend is not None:
                chart.legend.position = "b"
                chart.legend.overlay = False
                chart.legend.txPr = chart_text()
            if chart.title is not None:
                chart.title.overlay = False
            for axis in (chart.x_axis, chart.y_axis):
                axis.txPr = chart_text()
                axis.txPr.bodyPr.rot = 0
                axis.txPr.bodyPr.horzOverflow = "overflow"
                axis.txPr.bodyPr.vertOverflow = "overflow"
                # openpyxl's NoneSet converts the valid OOXML literal "none"
                # to Python None. Preserve it explicitly to prevent Excel
                # wrapping a Wafer caption into the width of one data point.
                axis.txPr.bodyPr.__dict__["wrap"] = "none"
                step = getattr(axis, "majorUnit", None)
                if step is not None:
                    decimals = max(0, int(-np.floor(np.log10(step))))
                    if not np.isclose(step / (10. ** -decimals), round(step / (10. ** -decimals))):
                        decimals += 1
                    axis.numFmt = "0.00E+00" if decimals > 4 else "0" + ("." + "0" * decimals if decimals else "")
                else:
                    axis.numFmt = "General"
                if axis.title is not None:
                    axis.title.overlay = False
                if axis.title and axis.title.tx and axis.title.tx.rich:
                    for paragraph in axis.title.tx.rich.p:
                        paragraph.pPr = ParagraphProperties(defRPr=properties)
            if chart.title and chart.title.tx and chart.title.tx.rich:
                rich = chart.title.tx.rich
                lines = "\n".join(run.t for paragraph in rich.p for run in paragraph.r).splitlines()
                if len(lines) > 1:
                    from openpyxl.drawing.text import RegularTextRun
                    rich.p = [Paragraph(r=[RegularTextRun(t=line)]) for line in lines]
                for index, paragraph in enumerate(rich.p):
                    title_properties = CharacterProperties(latin=DrawingFont(typeface=EXPORT_FONT),
                        sz=1400 if index == 0 else 1100, b=index == 0)
                    paragraph.pPr = ParagraphProperties(defRPr=title_properties)
                    for run in paragraph.r:
                        run.rPr = title_properties


def _embed_figure(sheet, figure, anchor, title, images):
    from .appearance import export_dpi, MAX_EXPORT_PIXELS
    buffer = BytesIO()
    figure.savefig(buffer, format="png", dpi=export_dpi(figure, 150, MAX_EXPORT_PIXELS),
                   facecolor="white", metadata={"Title": title})
    buffer.seek(0)
    picture = Image(buffer)
    picture.width = figure.get_figwidth() * 96
    picture.height = picture.width * figure.get_figheight() / figure.get_figwidth()
    sheet.add_image(picture, anchor)
    if images is not None:
        images.append((title, buffer.getvalue()))


def write_fullmap(sheet, frame, parameters, *, native=True, images=None, plot_labels=None, pictures=()):
    """FullMap's own measurements and mapped Trends; software Map/Radius PNGs."""
    from .analysis_export import _series, plot_figure
    from .measurements import wafer_identity_labels
    _write_frame(sheet, frame)
    labels = plot_labels or {"pad": True, "lot": True}
    captions = wafer_identity_labels(frame, [[row] for row in range(len(frame))],
                                    show_pad=labels.get("pad", True), show_lot=labels.get("lot", True)) if len(frame) else []
    layout = sequence_layout(captions)
    mapped = list(dict.fromkeys(parameter for parameter in parameters if parameter in frame))
    plot_column = len(frame.columns) + 2
    anchor_column = plot_column + len(mapped) + 3 if len(frame) else plot_column
    from .match_groups import normalized, wafer_identity_keys
    from openpyxl.styles import Alignment
    die = next((column for column in frame if normalized(column) in ("dieseq", "dieid", "diesequence")), None)
    dies = frame[die].astype(str).tolist() if die is not None else [""] * len(frame)
    for offset, header in enumerate(["Wafer", "Die Seq", *mapped]):
        sheet.cell(1, plot_column + offset, header).font = Font(name=EXPORT_FONT, bold=True)
        sheet.column_dimensions[get_column_letter(plot_column + offset)].width = 30 if offset == 0 else 16
    for index, (caption, die_label) in enumerate(zip(captions, dies), 2):
        sheet.cell(index, plot_column, caption).data_type = "s"
        sheet.cell(index, plot_column + 1, die_label).data_type = "s"
    keys = wafer_identity_keys(frame)
    for start, end, _key in _label_spans(keys if keys is not None else captions, limit=None):
        if start != end:
            sheet.merge_cells(start_row=start + 2, end_row=end + 2, start_column=plot_column, end_column=plot_column)
        sheet.cell(start + 2, plot_column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    def software_pictures(kind, column):
        cursor, width = 1, 0
        for title, contents, _note in pictures:
            if not title.startswith(kind):
                continue
            picture = Image(BytesIO(contents))
            ratio = min(1, 1200 / picture.width)
            picture.width, picture.height = picture.width * ratio, picture.height * ratio
            sheet.add_image(picture, f"{get_column_letter(column)}{cursor}")
            cursor += int(np.ceil(picture.height / 20)) + 3
            width = max(width, picture.width)
            if images is not None:
                images.append((sheet.title + " " + title, contents))
        return _next_plot_column(sheet, column, width) if width else column
    anchor_column = software_pictures("Wafer Map", anchor_column)
    anchor = get_column_letter(anchor_column)
    # Wafer Map, Trend and Radius each start beside the same data at row 1.
    software_pictures("Radius", _next_plot_column(sheet, anchor_column,
        layout["native_width" if native else "width"] * 96 / 2.54))
    cursor = 1
    if frame.empty:
        return
    categories = native_categories(sheet, 2, len(frame) + 1, [(plot_column, captions), (plot_column + 1, dies)])
    for slot, parameter in enumerate(mapped, plot_column + 2):
        values = pd.to_numeric(frame[parameter], errors="coerce").to_numpy(float)
        for index, value in enumerate(values):
            source = f"{get_column_letter(frame.columns.get_loc(parameter) + 1)}{index + 2}"
            sheet.cell(row=index + 2, column=slot,
                        value=f'=IFERROR(IF({source}="",NA(),{source}*1),NA())')
        if native:
            chart = LineChart()
            chart._identity_count = len(frame)
            chart.title = chart.y_axis.title = parameter
            chart.width, chart.height = layout["native_width"], PANEL_HEIGHT_CM
            chart.visible_cells_only = False
            chart.add_data(Reference(sheet, min_col=slot, min_row=1, max_row=len(frame) + 1), titles_from_data=True)
            chart.series[0].cat = categories
            chart_limits(chart, values)
            sheet.add_chart(chart, f"{anchor}{cursor}")
        else:
            plot = {"title": parameter, "x": "", "y": parameter,
                    "labels": captions, "series": [_series(parameter, np.arange(1, len(frame) + 1), values)]}
            _embed_figure(sheet, plot_figure(plot), f"{anchor}{cursor}", sheet.title + " " + parameter, images)
        cursor += 30


def save_atomic(book, path, *, cancelled=lambda: False):
    """Publish only a complete xlsx; a failed export leaves an existing file intact."""
    target = Path(path)
    handle = tempfile.NamedTemporaryFile(dir=target.parent, suffix=".xlsx", prefix=".export-", delete=False)
    temporary = Path(handle.name)
    handle.close()
    try:
        if cancelled():
            raise InterruptedError("Export cancelled")
        book.save(temporary)
        if cancelled():
            raise InterruptedError("Export cancelled")
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def set_export_font(figure):
    from matplotlib.text import Text
    for text in figure.findobj(match=Text):
        text.set_fontfamily(EXPORT_FONT)


def parameter_figure(parameter, data, card, stage, scope, *, panel="Match", reference_label="REF", unit="", sequence=None):
    """Matplotlib report from the same software values used in the Excel areas."""
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    width = MATCH_WIDTH_CM if panel == "Match" else TREND_WIDTH_CM if sequence is None else sequence["width"]
    figure = Figure(figsize=(width / 2.54, PANEL_HEIGHT_CM / 2.54), facecolor="white")
    FigureCanvasAgg(figure)
    axis = figure.subplots()
    raw, reference = data["Raw"].to_numpy(float), data["Reference"].to_numpy(float)
    values = (raw if stage.lower() == "final" else
              card.slope * raw + card.intercept if card is not None else np.full(len(raw), np.nan))
    valid = np.isfinite(raw) & np.isfinite(reference)
    if panel == "Match":
        style_plot_axis(axis, title=parameter, xlabel="PMISH", ylabel=reference_label)
        axis.scatter(raw[valid], reference[valid], s=12, color="#3776B6", alpha=.8, edgecolors="none")
        if valid.any() and card is not None and np.isfinite(card.slope) and np.isfinite(card.intercept):
            ends = np.array([raw[valid].min(), raw[valid].max()])
            axis.plot(ends, card.slope * ends + card.intercept, color="#D97732", linewidth=1.3)
            axis.text(.04, .97, f"y = {card.slope:.6g}x {card.intercept:+.6g}\nR² = {card.r_squared:.6g}",
                transform=axis.transAxes, va="top", ha="left", fontsize=10, color="#334155",
                bbox={"facecolor": "white", "edgecolor": "none", "alpha": .85, "pad": 4})
        if valid.any():
            plot_axis_limits(axis, reference[valid], raw[valid])
        bottom = 1.5
    else:
        rows = np.arange(1, len(data) + 1) if sequence is None else sequence["x"]
        ylabel = parameter if panel == "Trend" else f"Bias ({unit})" if panel == "Bias" and unit and unit != "1" else panel
        style_plot_axis(axis, title=parameter if panel == "Trend" else panel, xlabel="", ylabel=ylabel)
        if panel == "Trend":
            axis.plot(rows, values, marker="o", markersize=2.3, color="#3776B6", linewidth=1,
                      label="PMISH" if stage.lower() == "final" else "PMISH+Card")
            axis.plot(rows, reference, marker="s", markersize=2.3, color="#D97732", linewidth=1, linestyle="--", label=reference_label)
            plot_legend(axis)
            plot_axis_limits(axis, np.concatenate((values, reference)))
        else:
            bias = values - reference
            if panel == "Bias %":
                percent = np.full(len(reference), np.nan)
                np.divide(bias * 100, reference, out=percent, where=np.isfinite(reference) & (reference != 0))
                bias = percent
            axis.plot(rows, bias, marker="o", markersize=2.3, color="#3776B6", linewidth=1)
            axis.axhline(0, color="#8A8F98", linewidth=.7, linestyle="--")
            plot_axis_limits(axis, bias)
        selected = ([(float((start + end) / 2 + 1), label) for start, end, label in _label_spans(data["Label"])]
                    if sequence is None and "Label" in data else [] if sequence is None else sequence["ticks"])
        axis.set_xticks([position for position, _label in selected], [label for _position, label in selected], fontsize=10)
        axis.set_xlim(.5, max(1, len(data)) + .5)
        if sequence is not None:
            axis.set_xlabel(sequence["name"], fontsize=12, labelpad=8)
        bottom = .8 + .45 * max((len(label.splitlines()) for _position, label in selected), default=1)
    figure.subplots_adjust(left=1.6 / width, right=1 - .45 / width, bottom=bottom / PANEL_HEIGHT_CM, top=.81)
    set_export_font(figure)
    return figure
