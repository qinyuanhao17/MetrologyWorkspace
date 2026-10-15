"""Report layout through the two public Workbook export commands."""
from tests import test_excel_export as fixtures
from pathlib import Path
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook


class ExcelReportLayoutTests(unittest.TestCase):
    # Reuse only the public fixture/lifecycle, not the older test cases.
    tearDown = fixtures.ExcelExportTests.tearDown
    frames = fixtures.ExcelExportTests.frames
    export = fixtures.ExcelExportTests.export

    def setUp(self):
        self.window = fixtures.MatchingWindow()
        self.window.match_type.setCurrentText("NOVA")

    def test_compact_match_table_in_both_renderers(self):
        reference, raw = self.frames()
        raw["Lot ID"], raw["PAD Name"] = "L1", "ARRAY"
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        before = self.window.workspace_snapshot()
        with tempfile.TemporaryDirectory() as folder:
            books = []
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"{native}.xlsx", native=native))
                books.append(book)
                self.assertEqual([s.title for s in book if s.sheet_state == "visible"],
                                 ["Original (WO Card)", "Match (WO Card)", "Wafer (WO Card)", "RSQ (WO Card)"])
                sheet = book["Match (WO Card)"]
                self.assertEqual(sheet["A1"].value, "Wafer ID")
                self.assertEqual(sheet["C2"].value, "1.000")
                headers = next(row for row in sheet.values if "PMISH+Card" in row)
                self.assertEqual(headers[:7], ("Wafer", "Die Seq", "PMISH", "PMISH+Card", "NOVA", "Bias", "Bias %"))
                header_row = next(cell.row for row in sheet for cell in row if cell.value == "PMISH+Card")
                self.assertIn("LOT: L1", sheet.cell(header_row + 1, 1).value)
                self.assertIn("PAD: ARRAY", sheet.cell(header_row + 1, 1).value)
                objects = sheet._charts if native else sheet._images
                self.assertTrue(objects)
                self.assertTrue(all(obj.anchor._from.row == header_row - 2 for obj in objects))
                self.assertGreater(objects[0].anchor._from.col, 6)
                values = [c.value for row in sheet for c in row]
                self.assertNotIn("Card source", values)
                self.assertNotIn("Source row", values)
            self.assertEqual(list(books[0]["Match (WO Card)"].values), list(books[1]["Match (WO Card)"].values))
        after = self.window.workspace_snapshot()
        self.assertEqual(before.states, after.states)
        for name in before.frames:
            pd.testing.assert_frame_equal(before.frames[name], after.frames[name])

    def test_final_table_omits_second_card_column(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.result_mode.setCurrentText("Final")
        self.window.set_raw_frame(raw.assign(CD=["3.000", "5.000", "7.000"]))
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            sheet = book["Match (W Card)"]
            headers = next(row for row in sheet.values if "Bias %" in row)
            self.assertEqual(headers[:6], ("Wafer", "Die Seq", "PMISH", "NOVA", "Bias", "Bias %"))
            self.assertNotIn("PMISH+Card", [c.value for row in sheet for c in row])
            row = next(c.row for cells in sheet for c in cells if c.value == "PMISH") + 1
            self.assertEqual(book["Original (W Card)"]["C2"].value, "3.000")
            self.assertIn("'Original (W Card)'!$C$2", sheet.cell(row, 3).value)

    def test_both_renderers_use_physical_panel_sizes_and_full_width_parameter_heading(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"sizes-{native}.xlsx", native=native))
                sheet = book["Match (WO Card)"]
                objects = sheet._charts if native else sheet._images
                self.assertEqual(len(objects), 3, "Match, Trend and Bias must be independently sized panels")
                dimensions = [(obj.anchor.ext.cx / 360000, obj.anchor.ext.cy / 360000) for obj in objects]
                for actual, expected in zip(dimensions, ((15, 11), (25, 11), (25, 11))):
                    self.assertAlmostEqual(actual[0], expected[0], places=3)
                    self.assertAlmostEqual(actual[1], expected[1], places=3)
                self.assertIn("A7:G7", {str(area) for area in sheet.merged_cells.ranges})
                self.assertEqual(sheet["A7"].value, "CD")
                self.assertTrue(all(obj.anchor._from.row == 6 for obj in objects))

    def test_group_overview_and_complete_wafer_detail_export_side_by_side_in_both_modes(self):
        raw = pd.DataFrame({"Wafer ID": [f"WAFER-{index:03}" for index in range(20) for _ in range(3)],
            "Lot ID": ["LOT-001"] * 60, "PAD Name": ["ARRAY"] * 60, "Die Seq": [1, 2, 3] * 20,
            "CD": [str(index + 1) for index in range(60)], "TestFlag": [0] * 30 + [1] * 30})
        self.window.set_reference_frame(pd.DataFrame({"CD Reference": [2 * index + 3 for index in range(60)]}))
        self.window.set_raw_frame(raw)
        self.window.group_controls.restore({"enabled": True, "head_names": {"0": "Head A", "1": "Head B"}}, raw[["TestFlag"]])
        self.window.group_controls.apply_button.click()
        with tempfile.TemporaryDirectory() as folder:
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"identities-{native}.xlsx", native=native))
                sheet = book["Match (WO Card)"]
                panels = sheet._charts if native else sheet._images
                self.assertEqual(len(panels), 5, "Match, Group Trend/Bias, then Wafer Trend/Bias")
                dimensions = [panel.anchor.ext.cx / 360000 for panel in panels]
                self.assertEqual(dimensions[0], 15)
                self.assertGreater(dimensions[3], dimensions[1])
                self.assertGreaterEqual(dimensions[1], 25)
                self.assertTrue(all(left.anchor._from.col < right.anchor._from.col for left, right in zip(panels, panels[1:])))
                if native:
                    categories = panels[3].series[0].cat.multiLvlStrRef
                    self.assertTrue(categories.f.startswith("'Match (WO Card)'!"))
                    captions = [point.v for point in categories.multiLvlStrCache.lvl[1].pt]
                    for index in range(20):
                        self.assertTrue(any(f"WAFER-{index:03}" in caption for caption in captions))
                    self.assertEqual(len(captions), 20)

    def test_fullmap_matplotlib_trend_keeps_tight_ordinal_limits(self):
        from metrology_app.analysis_export import plot_figure
        plot = {"title": "DP", "x": "", "y": "DP", "labels": ["W1"] * 51,
                "series": [{"name": "DP", "line": True, "data": pd.DataFrame({"X": range(1, 52), "Y": range(51)})}]}
        figure = plot_figure(plot)
        self.assertEqual(figure.axes[0].get_xlim(), (.5, 51.5))
        self.assertAlmostEqual(figure.get_figwidth() * 2.54, 25, places=3)
        self.assertAlmostEqual(figure.get_figheight() * 2.54, 11, places=3)
        figure.clear()

    def test_native_trend_uses_editable_local_ranges_square_excel_style_and_merged_wafer(self):
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            path = self.window.export_excel(Path(folder) / "editable.xlsx")
            sheet = load_workbook(path)["Match (WO Card)"]
            trend = sheet._charts[1]
            category = trend.series[0].cat.multiLvlStrRef
            self.assertIsNotNone(category, "Native categories must reference the visible identity columns")
            self.assertEqual(category.f, "'Match (WO Card)'!$A$9:$B$11")
            self.assertEqual(trend.series[0].val.numRef.f, "'Match (WO Card)'!$D$9:$D$11")
            self.assertEqual(trend.series[1].val.numRef.f, "'Match (WO Card)'!$E$9:$E$11")
            self.assertEqual(trend.series[1].graphicalProperties.line.solidFill.srgbClr, "ED7D31")
            self.assertIn("A9:A11", {str(area) for area in sheet.merged_cells.ranges})
            self.assertEqual([sheet.cell(row, 1).value for row in range(9, 12)], ["W1", None, None])
            self.assertEqual([sheet.cell(row, 1).value for row in range(2, 5)], ["W1"] * 3)
            with ZipFile(path) as archive:
                chart = ET.fromstring(archive.read("xl/charts/chart2.xml"))
                ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                self.assertEqual(chart.find("c:roundedCorners", ns).get("val"), "0")

    def test_grouped_exports_reference_original_rows_instead_of_copying_reordered_values(self):
        raw = pd.DataFrame({"Wafer ID": ["W1", "W2", "W1", "W2"], "Die Seq": ["01", "01", "02", "02"],
                            "CD": ["1.000", "8.000", "2.000", "9.000"]})
        self.window.set_reference_frame(pd.DataFrame({"CD Reference": [3, 17, 5, 19]}))
        self.window.set_raw_frame(raw)
        self.window.group_controls.restore({"enabled": True, "trend_order": "original",
            "head_names": {"0": "Head A", "1": "Head B"}}, pd.DataFrame({"TestFlag": [0, 1, 0, 1]}))
        self.window.group_controls.apply_button.click()
        before = self.window.workspace_snapshot()
        with tempfile.TemporaryDirectory() as folder:
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"original-{native}.xlsx", native=native))
                original = book["Original (WO Card)"]
                self.assertEqual([original.cell(row, 3).value for row in range(2, 6)], ["1.000", "8.000", "2.000", "9.000"])
                sheet = book["Match (WO Card)"]
                header = next(cell.row for row in sheet for cell in row if cell.value == "PMISH+Card")
                self.assertEqual([sheet.cell(row, 1).value for row in range(header + 1, header + 5)], ["Head A", "Head A", "Head B", "Head B"])
                values = [sheet.cell(row, 4).value for row in range(header + 1, header + 5)]
                for formula, source_row in zip(values, (2, 4, 3, 5)):
                    self.assertIn(f"'Original (WO Card)'!$C${source_row}", formula)
                group = book["Match Head A (WO Card)"]
                self.assertIn("'Original (WO Card)'!$C$4", group["D4"].value)
        after = self.window.workspace_snapshot()
        self.assertEqual(before.states, after.states)
        for name in before.frames:
            pd.testing.assert_frame_equal(before.frames[name], after.frames[name])

    def test_wafer_sheet_references_original_in_first_appearance_order_in_both_modes(self):
        raw = pd.DataFrame({"Wafer ID": ["W2", "W1", "W2", "W1"], "Die Seq": ["01", "01", "02", "02"],
                            "CD": ["1.000", "8.000", "2.000", "9.000"]})
        self.window.set_reference_frame(pd.DataFrame({"CD Reference": [3, 17, 5, 19]}))
        self.window.set_raw_frame(raw)
        self.window.group_controls.restore({"enabled": True, "trend_order": "original",
            "head_names": {"0": "Head A", "1": "Head B"}}, pd.DataFrame({"TestFlag": [0, 1, 1, 0]}))
        self.window.group_controls.apply_button.click()
        before = self.window.workspace_snapshot()
        with tempfile.TemporaryDirectory() as folder:
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"wafer-{native}.xlsx", native=native))
                original = book["Original (WO Card)"]
                self.assertEqual((len(original._charts), len(original._images)), (0, 0))
                self.assertEqual([original.cell(row, 3).value for row in range(2, 6)], ["1.000", "8.000", "2.000", "9.000"])
                self.assertEqual(original["E1"].value, "CD Reference (NOVA)")
                self.assertEqual([original.cell(row, 5).value for row in range(2, 6)], [3, 17, 5, 19])
                wafer = book["Wafer (WO Card)"]
                self.assertEqual([wafer.cell(row, 3).value for row in range(3, 7)], ["01", "02", "01", "02"])
                self.assertEqual([wafer.cell(row, 2).value for row in range(3, 7)], ["W2", None, "W1", None])
                for row, source_row in zip(range(3, 7), (2, 4, 3, 5)):
                    self.assertIn(f"'Original (WO Card)'!$C${source_row}", wafer.cell(row, 4).value)
                    self.assertIn(f"'Original (WO Card)'!$E${source_row}", wafer.cell(row, 6).value)
                panels = wafer._charts if native else wafer._images
                self.assertEqual(len(panels), 3, "Match, Wafer Trend and Wafer Bias, without a Group overview")
                self.assertTrue(all(panel.anchor._from.row == 0 for panel in panels))
                self.assertGreaterEqual(panels[1].anchor.ext.cx / 360000, 25)
                if native:
                    categories = panels[1].series[0].cat.multiLvlStrRef
                    self.assertEqual(categories.f, "'Wafer (WO Card)'!$B$3:$C$6")
                    self.assertEqual([point.v for point in categories.multiLvlStrCache.lvl[1].pt], ["W2", "W1"])
        after = self.window.workspace_snapshot()
        self.assertEqual(before.states, after.states)
        for name in before.frames:
            pd.testing.assert_frame_equal(before.frames[name], after.frames[name])

    def test_fullmap_trend_uses_own_data_and_only_mapped_parameters(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        full = raw.assign(CD=["11.000", "21.000", "31.000"], Other=[101, 102, 103])
        self.window.set_preview_frame(full)
        self.window.run_analysis()
        expected = self.window.workbook.stage_frame("preview")["CD"].astype(float).tolist()
        with tempfile.TemporaryDirectory() as folder:
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"full-{native}.xlsx", native=native))
                sheet = book["FullMap (WO Card)"]
                self.assertEqual([sheet.cell(row, 3).value for row in range(2, 5)], expected)
                objects = sheet._charts if native else sheet._images
                self.assertEqual(len(objects), 1)
                self.assertEqual(objects[0].anchor._from.row, 0)
                self.assertGreater(objects[0].anchor._from.col, len(full.columns) - 1)
                if native:
                    source = objects[0].series[0].val.numRef.f
                    self.assertIn("'FullMap (WO Card)'", source)
                    formulas = [cell.value for row in sheet for cell in row
                                if cell.data_type == "f" and 'IF(C2=""' in cell.value]
                    self.assertEqual(len(formulas), 1)
                    self.assertIn("C2*1", formulas[0])

    def test_fullmap_reuses_only_mapped_drawn_panels_without_changing_child(self):
        from PyQt6.QtCore import Qt
        from PyQt6.QtTest import QTest
        import time
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        full = pd.DataFrame({"Wafer ID": ["OWN"] * 4, "Die Seq": [1, 2, 3, 4],
            "FIELD X": [-1, 1, 0, 0], "FIELD Y": [0, 0, -1, 1],
            "CD": [11, 21, 31, 41], "Other": [101, 102, 103, 104]})
        self.window.set_preview_frame(full)
        self.window.run_analysis()
        child = self.window.open_stage_workspace("preview")
        for index in range(child.parameter_list.topLevelItemCount()):
            item = child.parameter_list.topLevelItem(index)
            if item.text(0) in {"CD", "Other"}:
                item.setCheckState(0, Qt.CheckState.Checked)
        child.plot_page.selector.selectAll()
        child.plot_page.draw_maps()
        deadline = time.monotonic() + 30
        while child.plot_page.worker is not None and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertIsNone(child.plot_page.worker)
        self.assertEqual(len(child.plot_page.artists), 2)
        child.radius_page.selector.selectAll()
        child.radius_page.draw_plot()
        self.assertTrue(child.radius_page.ready)
        before = child.workspace_snapshot()
        result = child.plot_page.result
        with tempfile.TemporaryDirectory() as folder:
            books = []
            for native in (True, False):
                book = load_workbook(self.window.export_excel(Path(folder) / f"drawn-{native}.xlsx", native=native))
                books.append(book)
                sheet = book["FullMap (WO Card)"]
                # One software Map and one Radius, plus only CD's Trend.
                self.assertEqual((len(sheet._charts), len(sheet._images)), (1, 2) if native else (0, 3))
                self.assertTrue(all(image.anchor._from.col > len(full.columns) - 1 for image in sheet._images))
                self.assertTrue(all(obj.anchor._from.row == 0 for obj in [*sheet._charts, *sheet._images]))
            for index in (0, 1):
                self.assertEqual(books[0]["FullMap (WO Card)"]._images[index]._data(),
                                 books[1]["FullMap (WO Card)"]._images[index]._data())
        self.assertIs(child.plot_page.result, result)
        after = child.workspace_snapshot()
        self.assertEqual(before.states, after.states)
        for name in before.frames:
            pd.testing.assert_frame_equal(before.frames[name], after.frames[name])


