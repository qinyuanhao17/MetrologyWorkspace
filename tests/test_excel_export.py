"""Native Excel charts through the public export seam."""
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, ScatterChart
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication

from metrology_app.matching_window import MatchingWindow
from metrology_app.dynamic_window import DynamicWindow
from metrology_app.correlation_window import CorrelationWindow
from metrology_app.window import MainWindow
from metrology_app.matching import ParameterMapping
from unittest.mock import patch
from PIL import Image as PillowImage


APP = QApplication.instance() or QApplication([])


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.window = MatchingWindow()

    def tearDown(self):
        self.window.document.force_close = True
        self.window.close()
        self.window.deleteLater()
        APP.processEvents()

    def frames(self):
        reference = pd.DataFrame({"CD Reference": [3.0, 5.0, 7.0]})
        raw = pd.DataFrame({"Wafer ID": ["W1"] * 3, "Die Seq": [1, 2, 3],
                            "CD": ["1.000", "2.000", "3.000"]})
        return reference, raw

    def export(self, folder, name="export.xlsx"):
        path = Path(folder) / name
        self.window.export_excel(path)
        return load_workbook(path)

    def test_export_writes_native_match_charts_bound_to_cells(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            self.assertEqual(book.sheetnames,
                             ["_Plot data", "Original (WO Card)", "Match (WO Card)", "Wafer (WO Card)", "RSQ (WO Card)"])
            self.assertEqual(book["_Plot data"].sheet_state, "hidden")
            # Source strings stay text in the data sheets.
            self.assertEqual(book["Match (WO Card)"]["C2"].value, "1.000")
            match = book["Match (WO Card)"]
            charts = match._charts
            self.assertEqual([type(chart) for chart in charts],
                             [ScatterChart, LineChart, LineChart])
            self.assertEqual(charts[0].anchor._from.col, 8)
            self.assertTrue(all(left.anchor._from.col < right.anchor._from.col
                                for left, right in zip(charts, charts[1:])))
            self.assertEqual([chart.anchor._from.row for chart in charts], [6, 6, 6])
            self.assertNotEqual(charts[0].x_axis.numFmt.formatCode, "General")
            points = charts[0].series[0]
            self.assertEqual(points.xVal.numRef.f, "'Match (WO Card)'!$C$9:$C$11")
            self.assertEqual(points.yVal.numRef.f, "'Match (WO Card)'!$E$9:$E$11")
            self.assertEqual(charts[0].series[1].yVal.numRef.f, "'_Plot data'!$D$7:$D$8")
            self.assertEqual(charts[1].series[0].val.numRef.f, "'Match (WO Card)'!$D$9:$D$11")
            self.assertEqual(charts[1].series[1].val.numRef.f, "'Match (WO Card)'!$E$9:$E$11")
            self.assertEqual(charts[2].series[0].val.numRef.f, "'Match (WO Card)'!$F$9:$F$11")
            # The plot area applies the software Card instead of refitting.
            helper = book["_Plot data"]
            self.assertEqual(helper["K10"].value, 2.0)
            self.assertEqual(helper["K11"].value, 1.0)
            self.assertEqual(helper["K12"].value, 1.0)
            self.assertEqual(match["D9"].value, "=IF(ISNUMBER(C9),'_Plot data'!$K$10*C9+'_Plot data'!$K$11,NA())")
            self.assertEqual(match["F9"].value, '=IF(AND(ISNUMBER(D9),ISNUMBER(E9)),D9-E9,NA())')
            self.assertIn("y = 2x + 1", helper["K20"].value)
            for chart in charts:
                self.assertFalse(chart.title.overlay)
                self.assertFalse(chart.varyColors)
                self.assertEqual(chart.x_axis.axPos, "b")
                self.assertFalse(chart.x_axis.delete)
                self.assertFalse(chart.y_axis.delete)
                self.assertEqual(chart.x_axis.txPr.bodyPr.rot, 0)
                self.assertTrue(all(series.smooth is False for series in chart.series))
                for axis in (chart.x_axis, chart.y_axis):
                    if axis.title is not None:
                        self.assertIs(axis.title.overlay, False)
            category_range = charts[1].series[0].cat.multiLvlStrRef.f
            self.assertEqual(category_range, "'Match (WO Card)'!$A$9:$B$11")
            self.assertEqual(charts[2].series[0].cat.multiLvlStrRef.f, category_range)
            self.assertTrue(all(chart.visible_cells_only is False for chart in charts))
            from zipfile import ZipFile
            from xml.etree import ElementTree as ET
            with ZipFile(Path(folder) / "export.xlsx") as archive:
                xml = ET.fromstring(archive.read("xl/charts/chart2.xml"))
                body = xml.find(".//{http://schemas.openxmlformats.org/drawingml/2006/chart}catAx/"
                                "{http://schemas.openxmlformats.org/drawingml/2006/chart}txPr/"
                                "{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
                self.assertEqual(body.get("wrap"), "none")

    def test_export_writes_one_match_sheet_per_group(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.group_controls.restore(
            {"enabled": True, "head_names": {"0": "Head A", "1": "Head B"}},
            pd.DataFrame({"TestFlag": [0, 0, 1]}),
        )
        self.window.group_controls.apply_button.click()
        self.assertTrue(self.window.result.group_plan.enabled)
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            self.assertEqual(book.sheetnames,
                             ["_Plot data", "Original (WO Card)", "Match (WO Card)", "Wafer (WO Card)",
                              "Match Head A (WO Card)", "Match Head B (WO Card)", "All Trend (WO Card)", "RSQ (WO Card)",
                              "Group Fits (WO Card)", "Order (WO Card)"])
            for title, rows in (("Match Head A", 2), ("Match Head B", 1)):
                sheet = book[title + " (WO Card)"]
                self.assertTrue(sheet._charts, f"{title} needs its own charts")
                self.assertEqual(sheet["A2"].value, "Group")
                filled = [sheet.cell(row=row, column=1).value
                          for row in range(3, 3 + rows)]
                self.assertEqual(filled, [title.removeprefix("Match ")] * rows)
                self.assertIsNone(sheet.cell(row=3 + rows, column=1).value)
            trend = book["All Trend (WO Card)"]
            self.assertEqual(trend["C3"].value, "Head A PMISH")
            self.assertEqual(trend["I3"].value, "Head B Bias")
            self.assertEqual(trend["L1"].value, 2.)
            self.assertEqual(trend["L2"].value, 1.)
            self.assertEqual(trend["C4"].value, '=IF(ISNUMBER(B4),$L$1*B4+$L$2,NA())')
            self.assertEqual(trend["G4"].value, '=IF(ISNUMBER(F4),$N$1*F4+$N$2,NA())')
            references = [chart.series[index].val.numRef.f
                          for chart in trend._charts for index in range(len(chart.series))]
            self.assertEqual(references, ["'All Trend (WO Card)'!$C$4:$C$5", "'All Trend (WO Card)'!$D$4:$D$5",
                                          "'All Trend (WO Card)'!$G$4", "'All Trend (WO Card)'!$H$4"])
            rsq = book["RSQ (WO Card)"]
            self.assertEqual(rsq["C1"].value, "CD R²")
            self.assertEqual(rsq["A2"].value, "W1")
            self.assertEqual(rsq["C2"].value, 1)
            r_squared_range = rsq._charts[0].series[0].val.numRef.f
            self.assertTrue(r_squared_range.startswith("'RSQ (WO Card)'!$C$"), r_squared_range)

    def test_both_stages_have_card_suffixes_and_english_times_font(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.result_mode.setCurrentText("Final")
        self.window.set_raw_frame(raw.assign(CD=["3.000", "5.000", "7.000"]))
        self.window.result_mode.setCurrentText("Preview")
        self.window.run_analysis()
        before = self.window.workspace_snapshot()
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            self.assertIn("Match (WO Card)", book.sheetnames)
            self.assertIn("Match (W Card)", book.sheetnames)
            self.assertEqual(book["Match (W Card)"]["C2"].value, "3.000")
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            self.assertEqual(cell.font.name, "Times New Roman")
            self.assertEqual(book["Match (WO Card)"]["A8"].value, "Wafer")
        after = self.window.workspace_snapshot()
        self.assertEqual(before.states, after.states)
        for name in before.frames:
            pd.testing.assert_frame_equal(before.frames[name], after.frames[name])

    def test_matplotlib_export_keeps_excel_data_and_embeds_real_pngs(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            native = self.export(folder)
            target = Path(folder) / "plots.xlsx"
            path = self.window.export_matplotlib(target)
            self.assertEqual(path, target)
            pictures = load_workbook(path)
            self.assertFalse(any(sheet._charts for sheet in pictures))
            self.assertTrue(pictures["Match (WO Card)"]._images)
            self.assertEqual(pictures["Match (WO Card)"]._images[0].anchor._from.col, 8)
            for title in native.sheetnames:
                self.assertEqual(list(native[title].values), list(pictures[title].values))
            self.assertEqual(list(Path(folder).rglob("*.png")), [])
            from io import BytesIO
            with PillowImage.open(BytesIO(pictures["Match (WO Card)"]._images[0]._data())) as picture:
                self.assertIn("Matplotlib", picture.info.get("Software", ""))

    def test_matplotlib_menu_exports_one_workbook_without_png_sidecars(self):
        from PyQt6.QtWidgets import QFileDialog
        from PyQt6.QtTest import QTest
        import time
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "plots.xlsx"
            with patch.object(QFileDialog, "getSaveFileName", return_value=(str(target), "")):
                self.window.images_action.trigger()
            deadline = time.monotonic() + 30
            while not self.window.images_action.isEnabled() and time.monotonic() < deadline:
                QTest.qWait(10)
            self.assertTrue(target.is_file())
            self.assertEqual({p.name for p in Path(folder).iterdir()}, {"plots.xlsx"})
            self.assertTrue(load_workbook(target)["Match (WO Card)"]._images)

    def test_native_charts_match_the_report_palette_labels_and_ranges(self):
        reference, raw = self.frames()
        self.window.match_type.setCurrentText("NOVA")
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            match, trend, bias = book["Match (WO Card)"]._charts
            self.assertEqual(trend.series[0].graphicalProperties.line.solidFill.srgbClr, "5B9BD5")
            self.assertEqual(trend.series[1].graphicalProperties.line.solidFill.srgbClr, "ED7D31")
            self.assertEqual(trend.series[1].tx.strRef.f, "'Match (WO Card)'!E8")
            self.assertAlmostEqual(trend.y_axis.scaling.min, 2.)
            self.assertAlmostEqual(trend.y_axis.scaling.max, 8.)
            self.assertEqual(trend.y_axis.majorUnit, 1.)
            self.assertAlmostEqual(match.x_axis.scaling.min, .5)
            self.assertAlmostEqual(match.x_axis.scaling.max, 3.5)
            self.assertIsNone(match.legend)

    def test_dynamic_exports_own_cycle_values_and_software_three_sigma(self):
        window = DynamicWindow()
        try:
            window.set_table(pd.DataFrame({"Wafer ID": ["OWN"] * 6,
                "Die Seq": [1, 2, 1, 2, 1, 2], "DP": ["1.000", "2.000", "3.000", "4.000", "5.000", "6.000"]}), "Independent")
            for index in range(window.parameter_list.topLevelItemCount()):
                item = window.parameter_list.topLevelItem(index)
                if item.text(0) == "DP":
                    item.setCheckState(0, Qt.CheckState.Checked)
            APP.processEvents()

            with tempfile.TemporaryDirectory() as folder:
                path = window.export_excel(Path(folder) / "dynamic.xlsx")
                book = load_workbook(path)
                self.assertEqual(book["Data"]["D2"].value, "1.000")
                self.assertIn("Dynamic", book.sheetnames)
                self.assertTrue(book["Dynamic"]._charts)
                numbers = [cell.value for row in book["Dynamic"] for cell in row if cell.data_type == "n"]
                self.assertIn(6., numbers)  # Each Die has sample SD=2, so software 3 Sigma=6.
        finally:
            window.document.force_close = True
            window.close()
            window.deleteLater()
            APP.processEvents()

    def test_correlation_and_trend_export_independent_sources_and_draw_choices(self):
        window = CorrelationWindow()
        try:
            raw = pd.DataFrame({"Wafer ID": ["OWN"] * 4, "Die Seq": [1, 2, 3, 4],
                                "DP": ["1.000", "2.000", "3.000", "4.000"],
                                "EW": ["2.000", "4.000", "6.000", "8.000"]})
            reference = pd.DataFrame({"DP Ref": [10, 20, 30, 40], "EW Ref": [20, 40, 60, 80]})
            window.set_sources(reference, raw, (ParameterMapping("Alias", "DP Ref", "DP"),), "Preview")
            for tree in (window.parameter_list, window.reference_parameter_list):
                for index in range(tree.topLevelItemCount()):
                    item = tree.topLevelItem(index)
                    if item.flags() & Qt.ItemFlag.ItemIsUserCheckable:
                        item.setCheckState(0, Qt.CheckState.Checked)
            for page in (window.correlation_page, window.sequence_page):
                page.selector.selectAll()
                page.draw_plot()
                self.assertTrue(page.ready)
            before = window.workspace_snapshot()
            with tempfile.TemporaryDirectory() as folder:
                native = load_workbook(window.export_excel(Path(folder) / "own.xlsx"))
                self.assertTrue(native["Correlation"]._charts)
                self.assertTrue(native["Trend"]._charts)
                self.assertEqual(native["Raw Data"]["C2"].value, "1.000")
                self.assertEqual(native["Ref Data"]["C1"].value, "DP Ref")
                table = [cell.value for row in native["Trend"] for cell in row]
                self.assertIn("Source row", table)
                self.assertNotIn("Alias", table)
                pictures = load_workbook(window.export_matplotlib(Path(folder) / "pictures"))
                self.assertFalse(any(sheet._charts for sheet in pictures))
                self.assertTrue(pictures["Trend"]._images)
            after = window.workspace_snapshot()
            self.assertEqual(before.states, after.states)
            for name in before.frames:
                pd.testing.assert_frame_equal(before.frames[name], after.frames[name])
        finally:
            window.document.force_close = True
            window.close()
            window.deleteLater()
            APP.processEvents()

    def test_radius_export_reuses_the_existing_png_in_both_modes(self):
        window = MainWindow()
        try:
            window.set_table(pd.DataFrame({"Wafer ID": ["OWN"] * 3, "Die Seq": [1, 2, 3],
                "FIELD X": [-1, 0, 1], "FIELD Y": [0, 1, 0], "DP": ["1.000", "2.000", "3.000"]}), "Independent")
            for index in range(window.parameter_list.topLevelItemCount()):
                item = window.parameter_list.topLevelItem(index)
                if item.text(0) == "DP":
                    item.setCheckState(0, Qt.CheckState.Checked)
            window.radius_page.selector.selectAll()
            window.radius_page.draw_plot()
            self.assertTrue(window.radius_page.ready)
            with tempfile.TemporaryDirectory() as folder:
                native = load_workbook(window.export_excel(Path(folder) / "radius.xlsx"))
                pictures = load_workbook(window.export_matplotlib(Path(folder) / "pictures"))
                self.assertFalse(native["Radius"]._charts)
                self.assertEqual(native["Radius"]._images[0]._data(), pictures["Radius"]._images[0]._data())
        finally:
            window.document.force_close = True
            window.close()
            window.deleteLater()
            APP.processEvents()

    def test_export_failure_preserves_existing_file_and_cleans_partial_output(self):
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "existing.xlsx"
            self.window.export_excel(path)
            before = path.read_bytes()
            with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(OSError, "disk full"):
                    self.window.export_excel(path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(Path(folder).glob(".export-*")), [])

    def test_missing_values_stay_blank_and_source_text_is_not_a_formula(self):
        reference, raw = self.frames()
        raw.loc[0, "CD"] = ""
        raw["Note"] = ["=not a formula", "00123", "literal"]
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            book = self.export(folder)
            self.assertEqual(book["Match (WO Card)"]["D2"].data_type, "s")
            self.assertEqual(book["Match (WO Card)"]["D2"].value, "=not a formula")
            match = book["Match (WO Card)"]
            self.assertIsNone(book["Original (WO Card)"]["C2"].value)
            self.assertEqual(match["C9"].value,
                             '=IF(\'Original (WO Card)\'!$C$2="",NA(),IFERROR(\'Original (WO Card)\'!$C$2*1,NA()))')
            self.assertEqual(match["D9"].value, "=IF(ISNUMBER(C9),'_Plot data'!$K$10*C9+'_Plot data'!$K$11,NA())")
            self.assertEqual(match._charts[1].display_blanks, "gap")

    def test_native_fit_line_uses_the_same_paired_range_as_the_matplotlib_match(self):
        reference, raw = self.frames()
        reference.loc[2, "CD Reference"] = float("nan")
        raw.loc[2, "CD"] = "999.000"
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            match = self.export(folder)["Match (WO Card)"]
            self.assertEqual((book := match.parent)["_Plot data"]["C7"].value, 1)
            self.assertEqual(book["_Plot data"]["C8"].value, 2)

    def test_native_trend_keeps_different_secondary_units_in_separate_charts(self):
        window = CorrelationWindow()
        try:
            window.set_table(pd.DataFrame({"Wafer ID": ["OWN"] * 4, "Die Seq": [1, 2, 3, 4],
                "P [nm]": [1, 2, 3, 4], "Q [um]": [.1, .2, .3, .4], "R [nm]": [5, 6, 7, 8]}), "Independent")
            for index in range(window.parameter_list.topLevelItemCount()):
                item = window.parameter_list.topLevelItem(index)
                if item.flags() & Qt.ItemFlag.ItemIsUserCheckable:
                    item.setCheckState(0, Qt.CheckState.Checked)
            page = window.sequence_page
            window.set_trend_axis_mode("dual")
            page.selector.selectAll()
            page.draw_plot()
            page.set_overlay("P [nm]", "Q [um]")
            page.set_overlay("P [nm]", "R [nm]")
            with tempfile.TemporaryDirectory() as folder:
                book = load_workbook(window.export_excel(Path(folder) / "units.xlsx"))
                charts = book["Trend"]._charts
                self.assertEqual(len(charts), 4)  # The two standalone panels remain, as in the software.
                for chart in charts[:2]:
                    self.assertEqual(len(chart._charts), 2)
                    self.assertEqual(len(chart._charts[1].series), 1)
        finally:
            window.document.force_close = True
            window.close()
            window.deleteLater()
            APP.processEvents()

    def test_wafer_map_export_reuses_the_current_png_without_interpolation(self):
        from PyQt6.QtTest import QTest
        import time
        window = MainWindow()
        try:
            window.set_table(pd.DataFrame({"Wafer ID": ["OWN"] * 4, "Die Seq": [1, 2, 3, 4],
                "FIELD X": [-1, 1, 0, 0], "FIELD Y": [0, 0, -1, 1], "DP": ["1.000", "2.000", "3.000", "4.000"]}), "Independent")
            for index in range(window.parameter_list.topLevelItemCount()):
                item = window.parameter_list.topLevelItem(index)
                if item.text(0) == "DP":
                    item.setCheckState(0, Qt.CheckState.Checked)
            page = window.plot_page
            page.selector.selectAll()
            page.draw_maps()
            deadline = time.monotonic() + 30
            while page.worker is not None and time.monotonic() < deadline:
                QTest.qWait(10)
            self.assertIsNone(page.worker)
            self.assertTrue(page.artists)
            before = page.result
            with tempfile.TemporaryDirectory() as folder:
                native = load_workbook(window.export_excel(Path(folder) / "maps.xlsx"))
                pictures = load_workbook(window.export_matplotlib(Path(folder) / "pictures"))
                self.assertEqual(native["Wafer Map"]._images[0]._data(), pictures["Wafer Map"]._images[0]._data())
                self.assertIs(page.result, before)
                self.assertIsNone(page.worker)
        finally:
            window.document.force_close = True
            window.close()
            window.deleteLater()
            APP.processEvents()

    def test_excel_menu_exports_in_background_and_reports_completion(self):
        from PyQt6.QtWidgets import QFileDialog
        from PyQt6.QtTest import QTest
        import time
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "background.xlsx"
            with patch.object(QFileDialog, "getSaveFileName", return_value=(str(target), "")):
                self.window.export_action.trigger()
            self.assertFalse(self.window.export_action.isEnabled())
            deadline = time.monotonic() + 30
            while not self.window.export_action.isEnabled() and time.monotonic() < deadline:
                QTest.qWait(10)
            self.assertTrue(self.window.export_action.isEnabled())
            self.assertIn("Match (WO Card)", load_workbook(target).sheetnames)

    def test_cancel_during_excel_write_preserves_the_previous_export(self):
        from threading import Event
        from PyQt6.QtWidgets import QFileDialog, QProgressDialog, QPushButton
        from PyQt6.QtTest import QTest
        from openpyxl import Workbook
        import time
        reference, raw = self.frames()
        self.window.set_reference_frame(reference)
        self.window.set_raw_frame(raw)
        self.window.run_analysis()
        entered, release = Event(), Event()
        original_save = Workbook.save
        def delayed_save(book, path):
            entered.set()
            if not release.wait(10):
                raise OSError("test write did not resume")
            original_save(book, path)
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "existing.xlsx"
            self.window.export_excel(target)
            before = target.read_bytes()
            with patch.object(Workbook, "save", delayed_save), patch.object(QFileDialog, "getSaveFileName", return_value=(str(target), "")):
                try:
                    self.window.export_action.trigger()
                    deadline = time.monotonic() + 10
                    while not entered.is_set() and time.monotonic() < deadline:
                        QTest.qWait(10)
                    self.assertTrue(entered.is_set())
                    progress = self.window.findChild(QProgressDialog)
                    next(button for button in progress.findChildren(QPushButton) if button.text() == "Cancel").click()
                finally:
                    release.set()
                while not self.window.export_action.isEnabled() and time.monotonic() < deadline:
                    QTest.qWait(10)
            self.assertTrue(self.window.export_action.isEnabled())
            self.assertEqual(target.read_bytes(), before)
            self.assertEqual(list(Path(folder).glob(".export-*")), [])


if __name__ == "__main__":
    unittest.main()
